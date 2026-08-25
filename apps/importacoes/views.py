from io import BytesIO
from decimal import Decimal, InvalidOperation
from difflib import get_close_matches
import unicodedata
import json
import re

from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.db.models import Sum
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

from apps.cadastros.models import Escola, Fornecedor, Item, Municipio, UnidadeMedida, validar_cnpj, validar_cpf

from .forms import ImportacaoPlanilhaForm
from apps.pregoes.models import Pregao, PregaoItem, QuantitativoEscola, QuantitativoPregao


CABECALHOS_FORNECEDORES = [
    "RAZÃO SOCIAL",
    "CNPJ",
    "ENDEREÇO",
    "TELEFONE",
    "REPRESENTANTE LEGAL",
    "RG REPRESENTANTE",
    "ÓRGÃO EXPEDIDOR",
    "CPF REPRESENTANTE",
    "E-MAIL",
    "NOME DO BANCO",
    "AGÊNCIA",
    "CONTA CORRENTE",
    "PIX",
    "ATIVO",
]


def usuario_pode_importar(user):
    if not user.is_authenticated:
        return False

    if user.is_superuser:
        return True

    perfil = getattr(user, "perfil_acesso", None)

    return bool(
        perfil
        and perfil.perfil == "administrador"
        and user.is_active
    )


def somente_digitos(valor):
    return "".join(ch for ch in str(valor or "") if ch.isdigit())


def texto_celula(valor):
    if valor is None:
        return ""

    return str(valor).strip()


def limpar_cnpj(valor):
    digitos = somente_digitos(valor)

    if len(digitos) == 14:
        return f"{digitos[:2]}.{digitos[2:5]}.{digitos[5:8]}/{digitos[8:12]}-{digitos[12:]}"

    return texto_celula(valor)


def limpar_cpf(valor):
    digitos = somente_digitos(valor)

    if len(digitos) == 11:
        return f"{digitos[:3]}.{digitos[3:6]}.{digitos[6:9]}-{digitos[9:]}"

    return texto_celula(valor)


def texto_ativo(valor):
    texto = texto_celula(valor).lower()

    if texto in ["", "sim", "s", "ativo", "1", "true", "verdadeiro"]:
        return True

    if texto in ["não", "nao", "n", "inativo", "0", "false", "falso"]:
        return False

    return True


def validar_email_opcional(email):
    email = texto_celula(email)

    if not email:
        return ""

    validate_email(email)
    return email


def cabecalhos_normalizados(linha):
    return [texto_celula(celula).upper() for celula in linha]


def validar_cabecalho_planilha(linha):
    encontrados = cabecalhos_normalizados(linha)

    return encontrados[:len(CABECALHOS_FORNECEDORES)] == CABECALHOS_FORNECEDORES


def montar_linha_fornecedor(numero_linha, valores, atualizar_existentes):
    dados = dict(zip(CABECALHOS_FORNECEDORES, valores))

    razao_social = texto_celula(dados.get("RAZÃO SOCIAL"))
    cnpj = limpar_cnpj(dados.get("CNPJ"))
    endereco = texto_celula(dados.get("ENDEREÇO"))
    telefone = texto_celula(dados.get("TELEFONE"))
    representante_legal = texto_celula(dados.get("REPRESENTANTE LEGAL"))
    rg_representante = texto_celula(dados.get("RG REPRESENTANTE"))
    orgao_expedidor = texto_celula(dados.get("ÓRGÃO EXPEDIDOR"))
    cpf_representante = limpar_cpf(dados.get("CPF REPRESENTANTE"))
    email = texto_celula(dados.get("E-MAIL"))
    nome_banco = texto_celula(dados.get("NOME DO BANCO"))
    agencia = texto_celula(dados.get("AGÊNCIA"))
    conta_corrente = texto_celula(dados.get("CONTA CORRENTE"))
    pix = texto_celula(dados.get("PIX"))
    ativo = texto_ativo(dados.get("ATIVO"))

    erros = []

    if not razao_social:
        erros.append("Razão social é obrigatória.")

    if not cnpj:
        erros.append("CNPJ é obrigatório.")
    else:
        try:
            validar_cnpj(cnpj)
        except ValidationError:
            erros.append("CNPJ inválido.")

    if not endereco:
        erros.append("Endereço é obrigatório.")

    if not representante_legal:
        erros.append("Representante legal é obrigatório.")

    if not rg_representante:
        erros.append("RG do representante é obrigatório.")

    if not orgao_expedidor:
        erros.append("Órgão expedidor é obrigatório.")

    if not cpf_representante:
        erros.append("CPF do representante é obrigatório.")
    else:
        try:
            validar_cpf(cpf_representante)
        except ValidationError:
            erros.append("CPF do representante inválido.")

    if email:
        try:
            email = validar_email_opcional(email)
        except ValidationError:
            erros.append("E-mail inválido.")

    fornecedor_existente = None

    if cnpj:
        fornecedor_existente = Fornecedor.objects.filter(
            cnpj__regex=rf"\\D*{somente_digitos(cnpj)}\\D*"
        ).first()

        if not fornecedor_existente:
            fornecedor_existente = Fornecedor.objects.filter(
                cnpj=cnpj,
            ).first()

    acao = "Criar"

    if fornecedor_existente:
        if atualizar_existentes:
            acao = "Atualizar"
        else:
            acao = "Ignorar"
            erros.append("Fornecedor já existe com este CNPJ.")

    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "dados": {
            "razao_social": razao_social,
            "cnpj": cnpj,
            "endereco": endereco,
            "telefone": telefone,
            "representante_legal": representante_legal,
            "rg_representante": rg_representante,
            "orgao_expedidor_representante": orgao_expedidor,
            "cpf_representante": cpf_representante,
            "email": email,
            "nome_banco": nome_banco,
            "agencia": agencia,
            "conta_corrente": conta_corrente,
            "pix": pix,
            "ativo": ativo,
        },
    }


def ler_planilha_fornecedores(arquivo, atualizar_existentes=True):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha os nomes das colunas."
        )

    linhas = []
    cnpjs_lidos = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_FORNECEDORES) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_fornecedor(
            numero_linha,
            valores,
            atualizar_existentes,
        )

        cnpj_digitos = somente_digitos(linha["dados"]["cnpj"])

        if cnpj_digitos:
            if cnpj_digitos in cnpjs_lidos:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append("CNPJ duplicado dentro da própria planilha.")
            else:
                cnpjs_lidos.add(cnpj_digitos)

        linhas.append(linha)

    return linhas


def criar_modelo_fornecedores_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Fornecedores"

    sheet.append(CABECALHOS_FORNECEDORES)
    sheet.append([
        "EMPRESA EXEMPLO LTDA",
        "00.000.000/0000-00",
        "Rua Exemplo, nº 100, Centro",
        "(66) 99999-9999",
        "NOME DO REPRESENTANTE",
        "1234567",
        "SSP/MT",
        "000.000.000-00",
        "fornecedor@email.com",
        "Banco do Brasil",
        "0000-0",
        "00000-0",
        "chave-pix@email.com",
        "Sim",
    ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    larguras = {
        "A": 34,
        "B": 20,
        "C": 38,
        "D": 18,
        "E": 30,
        "F": 18,
        "G": 18,
        "H": 20,
        "I": 28,
        "J": 22,
        "K": 14,
        "L": 18,
        "M": 28,
        "N": 12,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    sheet.row_dimensions[1].height = 32

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_fornecedores(request):
    workbook = criar_modelo_fornecedores_workbook()
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_fornecedores.xlsx"'

    return response


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_fornecedores(request):
    linhas = []
    resumo = None

    if request.method == "POST":
        acao = request.POST.get("acao")

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]
                atualizar_existentes = form.cleaned_data["atualizar_existentes"]

                try:
                    linhas = ler_planilha_fornecedores(
                        arquivo,
                        atualizar_existentes=atualizar_existentes,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [linha for linha in linhas if linha["valido"]]
                    erros = [linha for linha in linhas if not linha["valido"]]

                    request.session["importacao_fornecedores_linhas"] = linhas
                    request.session["importacao_fornecedores_atualizar"] = atualizar_existentes

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                        "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija as linhas inválidas antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )
        elif acao == "confirmar":
            linhas = request.session.get("importacao_fornecedores_linhas", [])

            if not linhas:
                messages.error(request, "Nenhuma importação validada foi encontrada. Envie a planilha novamente.")
                return redirect("importacoes:fornecedores")

            linhas_invalidas = [linha for linha in linhas if not linha.get("valido")]

            if linhas_invalidas:
                messages.error(request, "Ainda existem linhas inválidas. Corrija a planilha e valide novamente.")
                return redirect("importacoes:fornecedores")

            criados = 0
            atualizados = 0

            for linha in linhas:
                dados = linha["dados"]
                cnpj_digitos = somente_digitos(dados["cnpj"])

                fornecedor = None

                for existente in Fornecedor.objects.all():
                    if somente_digitos(existente.cnpj) == cnpj_digitos:
                        fornecedor = existente
                        break

                if fornecedor:
                    for campo, valor in dados.items():
                        setattr(fornecedor, campo, valor)
                    fornecedor.save()
                    atualizados += 1
                else:
                    Fornecedor.objects.create(**dados)
                    criados += 1

            request.session.pop("importacao_fornecedores_linhas", None)
            request.session.pop("importacao_fornecedores_atualizar", None)

            messages.success(
                request,
                f"Importação concluída. Fornecedores criados: {criados}. Fornecedores atualizados: {atualizados}.",
            )
            return redirect("cadastros:fornecedores")
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/fornecedores.html",
        {
            "form": form,
            "linhas": linhas,
            "resumo": resumo,
        },
    )


CABECALHOS_ESCOLAS = [
    "NOME DA ESCOLA",
    "CNPJ",
    "ENDEREÇO",
    "NÚMERO",
    "BAIRRO",
    "MUNICÍPIO",
    "UF",
    "PRESIDENTE DO CDCE",
    "RG PRESIDENTE",
    "CPF PRESIDENTE",
    "ATIVO",
]


def validar_cabecalho_planilha_escolas(linha):
    encontrados = cabecalhos_normalizados(linha)

    return encontrados[:len(CABECALHOS_ESCOLAS)] == CABECALHOS_ESCOLAS


def buscar_escola_por_cnpj(cnpj):
    cnpj_digitos = somente_digitos(cnpj)

    if not cnpj_digitos:
        return None

    for escola in Escola.objects.all():
        if somente_digitos(escola.cnpj) == cnpj_digitos:
            return escola

    return None


def normalizar_texto_comparacao(texto):
    texto = texto_celula(texto).lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(ch for ch in texto if not unicodedata.combining(ch))
    texto = re.sub(r"[^a-z0-9]+", " ", texto)
    texto = re.sub(r"\s+", " ", texto).strip()

    return texto


def buscar_municipio_cadastrado(nome, uf):
    nome = texto_celula(nome)
    uf = texto_celula(uf).upper() or "MT"

    if not nome:
        return None

    # Primeiro tenta comparação exata ignorando maiúsculas/minúsculas.
    municipio = Municipio.objects.filter(
        nome__iexact=nome,
        uf__iexact=uf,
    ).first()

    if municipio:
        return municipio

    # Depois tenta comparação normalizada, ignorando acentos, pontuação e espaços.
    nome_normalizado = normalizar_texto_comparacao(nome)

    for municipio in Municipio.objects.filter(uf__iexact=uf):
        if normalizar_texto_comparacao(municipio.nome) == nome_normalizado:
            return municipio

    return None


def sugerir_municipios(nome, uf, limite=3):
    nome_normalizado = normalizar_texto_comparacao(nome)
    uf = texto_celula(uf).upper() or "MT"

    municipios = list(Municipio.objects.filter(uf__iexact=uf).order_by("nome"))

    mapa_normalizado = {
        normalizar_texto_comparacao(municipio.nome): municipio.nome
        for municipio in municipios
    }

    sugestoes_normalizadas = get_close_matches(
        nome_normalizado,
        list(mapa_normalizado.keys()),
        n=limite,
        cutoff=0.72,
    )

    return [mapa_normalizado[item] for item in sugestoes_normalizadas]



def opcoes_municipio_para_correcao(nome, uf):
    """
    Retorna municípios cadastrados para o usuário selecionar na tela de prévia.

    A lista prioriza os municípios mais parecidos com o texto da planilha,
    mas também inclui os demais municípios da mesma UF para permitir correção manual.
    """
    uf = texto_celula(uf).upper() or "MT"
    municipios = list(Municipio.objects.filter(uf__iexact=uf).order_by("nome"))

    sugestoes_nomes = sugerir_municipios(nome, uf, limite=5)
    sugestoes_normalizadas = {
        normalizar_texto_comparacao(nome_sugerido)
        for nome_sugerido in sugestoes_nomes
    }

    municipios_ordenados = sorted(
        municipios,
        key=lambda municipio: (
            0 if normalizar_texto_comparacao(municipio.nome) in sugestoes_normalizadas else 1,
            municipio.nome,
        ),
    )

    return [
        {
            "id": municipio.id,
            "nome": municipio.nome,
            "uf": municipio.uf,
        }
        for municipio in municipios_ordenados
    ]


def obter_municipio_por_correcao(request, numero_linha):
    municipio_id = request.POST.get(f"municipio_corrigido_{numero_linha}")

    if not municipio_id:
        return None

    return Municipio.objects.filter(id=municipio_id).first()


def montar_linha_escola(numero_linha, valores, atualizar_existentes):
    dados = dict(zip(CABECALHOS_ESCOLAS, valores))

    nome_escola = texto_celula(dados.get("NOME DA ESCOLA"))
    cnpj = limpar_cnpj(dados.get("CNPJ"))
    endereco = texto_celula(dados.get("ENDEREÇO"))
    numero = texto_celula(dados.get("NÚMERO"))
    bairro = texto_celula(dados.get("BAIRRO"))
    municipio_nome = texto_celula(dados.get("MUNICÍPIO"))
    uf = texto_celula(dados.get("UF")).upper() or "MT"
    presidente_cdce = texto_celula(dados.get("PRESIDENTE DO CDCE"))
    rg_presidente = texto_celula(dados.get("RG PRESIDENTE"))
    cpf_presidente = limpar_cpf(dados.get("CPF PRESIDENTE"))
    ativo = texto_ativo(dados.get("ATIVO"))

    erros = []

    if not nome_escola:
        erros.append("Nome da escola é obrigatório.")

    if not cnpj:
        erros.append("CNPJ é obrigatório.")
    else:
        try:
            validar_cnpj(cnpj)
        except ValidationError:
            erros.append("CNPJ inválido.")

    if not endereco:
        erros.append("Endereço é obrigatório.")

    if not municipio_nome:
        erros.append("Município é obrigatório.")

    if not uf:
        erros.append("UF é obrigatória.")
    elif len(uf) != 2:
        erros.append("UF deve conter 2 letras.")

    municipio_encontrado = None

    erro_municipio = False
    municipio_opcoes = []

    if municipio_nome and uf and len(uf) == 2:
        municipio_encontrado = buscar_municipio_cadastrado(municipio_nome, uf)

        if not municipio_encontrado:
            erro_municipio = True
            municipio_opcoes = opcoes_municipio_para_correcao(municipio_nome, uf)
            sugestoes = sugerir_municipios(municipio_nome, uf)

            if sugestoes:
                erros.append(
                    "Município não encontrado no sistema. "
                    f"Selecione o município correto na coluna Correção. Sugestões: {', '.join(sugestoes)}."
                )
            else:
                erros.append(
                    "Município não encontrado no sistema. "
                    "Selecione o município correto na coluna Correção ou cadastre o município antes de importar."
                )

    if not presidente_cdce:
        erros.append("Presidente do CDCE é obrigatório.")

    if not rg_presidente:
        erros.append("RG do presidente é obrigatório.")

    if not cpf_presidente:
        erros.append("CPF do presidente é obrigatório.")
    else:
        try:
            validar_cpf(cpf_presidente)
        except ValidationError:
            erros.append("CPF do presidente inválido.")

    escola_existente = buscar_escola_por_cnpj(cnpj)

    acao = "Criar"

    if escola_existente:
        if atualizar_existentes:
            acao = "Atualizar"
        else:
            acao = "Ignorar"
            erros.append("Escola já existe com este CNPJ.")

    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "erro_municipio": erro_municipio,
        "municipio_opcoes": municipio_opcoes,
        "dados": {
            "nome_escola": nome_escola,
            "cnpj": cnpj,
            "endereco": endereco,
            "numero": numero,
            "bairro": bairro,
            "municipio_nome": municipio_nome,
            "uf": uf,
            "presidente_cdce": presidente_cdce,
            "rg_presidente": rg_presidente,
            "cpf_presidente": cpf_presidente,
            "ativo": ativo,
            "municipio_corrigido_id": "",
            "municipio_corrigido_nome": "",
            "municipio_corrigido_uf": "",
        },
    }


def ler_planilha_escolas(arquivo, atualizar_existentes=True):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha_escolas(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha os nomes das colunas."
        )

    linhas = []
    cnpjs_lidos = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_ESCOLAS) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_escola(
            numero_linha,
            valores,
            atualizar_existentes,
        )

        cnpj_digitos = somente_digitos(linha["dados"]["cnpj"])

        if cnpj_digitos:
            if cnpj_digitos in cnpjs_lidos:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append("CNPJ duplicado dentro da própria planilha.")
            else:
                cnpjs_lidos.add(cnpj_digitos)

        linhas.append(linha)

    return linhas


def criar_modelo_escolas_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Escolas"

    sheet.append(CABECALHOS_ESCOLAS)
    sheet.append([
        "EE EXEMPLO",
        "00.000.000/0000-00",
        "Rua Exemplo",
        "100",
        "Centro",
        "Sinop",
        "MT",
        "NOME DO PRESIDENTE",
        "1234567",
        "000.000.000-00",
        "Sim",
    ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    larguras = {
        "A": 34,
        "B": 20,
        "C": 34,
        "D": 12,
        "E": 20,
        "F": 22,
        "G": 10,
        "H": 30,
        "I": 18,
        "J": 20,
        "K": 12,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    sheet.row_dimensions[1].height = 32

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_escolas(request):
    workbook = criar_modelo_escolas_workbook()
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_escolas.xlsx"'

    return response


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_escolas(request):
    linhas = []
    resumo = None

    if request.method == "POST":
        acao = request.POST.get("acao")

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]
                atualizar_existentes = form.cleaned_data["atualizar_existentes"]

                try:
                    linhas = ler_planilha_escolas(
                        arquivo,
                        atualizar_existentes=atualizar_existentes,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [linha for linha in linhas if linha["valido"]]
                    erros = [linha for linha in linhas if not linha["valido"]]

                    request.session["importacao_escolas_linhas"] = linhas
                    request.session["importacao_escolas_atualizar"] = atualizar_existentes

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                        "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija as linhas inválidas antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )

        elif acao == "confirmar":
            linhas = request.session.get("importacao_escolas_linhas", [])

            if not linhas:
                messages.error(request, "Nenhuma importação validada foi encontrada. Envie a planilha novamente.")
                return redirect("importacoes:escolas")

            # Permite corrigir município diretamente na tela de prévia,
            # sem precisar editar e reenviar a planilha.
            for linha in linhas:
                if linha.get("erro_municipio"):
                    municipio_corrigido = obter_municipio_por_correcao(request, linha["linha"])

                    if municipio_corrigido:
                        erros_restantes = [
                            erro for erro in linha.get("erros", [])
                            if "Município não encontrado" not in erro
                        ]

                        linha["erros"] = erros_restantes
                        linha["erro_municipio"] = False
                        linha["dados"]["municipio_corrigido_id"] = municipio_corrigido.id
                        linha["dados"]["municipio_corrigido_nome"] = municipio_corrigido.nome
                        linha["dados"]["municipio_corrigido_uf"] = municipio_corrigido.uf
                        linha["dados"]["municipio_nome"] = municipio_corrigido.nome
                        linha["dados"]["uf"] = municipio_corrigido.uf

                        if not erros_restantes:
                            linha["valido"] = True
                            if linha.get("acao") == "Erro":
                                linha["acao"] = "Atualizar" if buscar_escola_por_cnpj(linha["dados"]["cnpj"]) else "Criar"

            linhas_invalidas = [linha for linha in linhas if not linha.get("valido")]

            if linhas_invalidas:
                request.session["importacao_escolas_linhas"] = linhas
                messages.error(
                    request,
                    "Ainda existem linhas inválidas. Se o erro for município, selecione o município correto na coluna Correção e confirme novamente.",
                )

                validos = [linha for linha in linhas if linha.get("valido")]
                erros = [linha for linha in linhas if not linha.get("valido")]
                resumo = {
                    "total": len(linhas),
                    "validos": len(validos),
                    "erros": len(erros),
                    "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                    "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                }

                form = ImportacaoPlanilhaForm()
                return render(
                    request,
                    "importacoes/escolas.html",
                    {
                        "form": form,
                        "linhas": linhas,
                        "resumo": resumo,
                    },
                )

            criadas = 0
            atualizadas = 0

            for linha in linhas:
                dados = linha["dados"]

                municipio = None

                if dados.get("municipio_corrigido_id"):
                    municipio = Municipio.objects.filter(id=dados["municipio_corrigido_id"]).first()

                if not municipio:
                    municipio = buscar_municipio_cadastrado(
                        dados["municipio_nome"],
                        dados["uf"],
                    )

                if not municipio:
                    messages.error(
                        request,
                        f"Município não encontrado para a escola {dados['nome_escola']}. "
                        "Selecione o município correto na prévia ou corrija a planilha.",
                    )
                    return redirect("importacoes:escolas")

                escola = buscar_escola_por_cnpj(dados["cnpj"])

                dados_escola = {
                    "nome_escola": dados["nome_escola"],
                    "cnpj": dados["cnpj"],
                    "endereco": dados["endereco"],
                    "numero": dados["numero"],
                    "bairro": dados["bairro"],
                    "municipio": municipio,
                    "presidente_cdce": dados["presidente_cdce"],
                    "rg_presidente": dados["rg_presidente"],
                    "cpf_presidente": dados["cpf_presidente"],
                    "ativo": dados["ativo"],
                }

                if escola:
                    for campo, valor in dados_escola.items():
                        setattr(escola, campo, valor)
                    escola.save()
                    atualizadas += 1
                else:
                    Escola.objects.create(**dados_escola)
                    criadas += 1

            request.session.pop("importacao_escolas_linhas", None)
            request.session.pop("importacao_escolas_atualizar", None)

            messages.success(
                request,
                f"Importação concluída. Escolas criadas: {criadas}. "
                f"Escolas atualizadas: {atualizadas}.",
            )
            return redirect("cadastros:escolas")
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/escolas.html",
        {
            "form": form,
            "linhas": linhas,
            "resumo": resumo,
        },
    )


CABECALHOS_ITENS = [
    "NOME DO ITEM",
    "UNIDADE DE MEDIDA",
    "DESCRIÇÃO",
    "ATIVO",
]


def validar_cabecalho_planilha_itens(linha):
    encontrados = cabecalhos_normalizados(linha)

    return encontrados[:len(CABECALHOS_ITENS)] == CABECALHOS_ITENS


def buscar_unidade_medida_cadastrada(texto_unidade):
    texto_unidade = texto_celula(texto_unidade)

    if not texto_unidade:
        return None

    texto_normalizado = normalizar_texto_comparacao(texto_unidade)

    # Primeiro tenta pelo código interno.
    unidade = UnidadeMedida.objects.filter(codigo__iexact=texto_unidade).first()
    if unidade:
        return unidade

    # Depois tenta por nome ou sigla, ignorando acentos e variações de caixa.
    for unidade in UnidadeMedida.objects.all().order_by("nome"):
        possibilidades = [
            unidade.codigo,
            unidade.nome,
            unidade.sigla,
            str(unidade),
        ]

        for possibilidade in possibilidades:
            if possibilidade and normalizar_texto_comparacao(possibilidade) == texto_normalizado:
                return unidade

    return None


def sugerir_unidades_medida(texto_unidade, limite=5):
    texto_normalizado = normalizar_texto_comparacao(texto_unidade)

    unidades = list(UnidadeMedida.objects.filter(ativo=True).order_by("nome"))

    mapa_normalizado = {}

    for unidade in unidades:
        rotulo = unidade.nome
        if unidade.sigla:
            rotulo = f"{unidade.nome} ({unidade.sigla})"

        for possibilidade in [unidade.codigo, unidade.nome, unidade.sigla, str(unidade)]:
            if possibilidade:
                mapa_normalizado[normalizar_texto_comparacao(possibilidade)] = rotulo

    sugestoes_normalizadas = get_close_matches(
        texto_normalizado,
        list(mapa_normalizado.keys()),
        n=limite,
        cutoff=0.60,
    )

    sugestoes = []

    for item in sugestoes_normalizadas:
        rotulo = mapa_normalizado[item]
        if rotulo not in sugestoes:
            sugestoes.append(rotulo)

    return sugestoes


def opcoes_unidade_para_correcao(texto_unidade):
    unidades = list(UnidadeMedida.objects.filter(ativo=True).order_by("nome"))

    sugestoes = sugerir_unidades_medida(texto_unidade, limite=5)
    sugestoes_normalizadas = {
        normalizar_texto_comparacao(sugestao)
        for sugestao in sugestoes
    }

    unidades_ordenadas = sorted(
        unidades,
        key=lambda unidade: (
            0 if normalizar_texto_comparacao(unidade.nome) in sugestoes_normalizadas
            or normalizar_texto_comparacao(str(unidade)) in sugestoes_normalizadas else 1,
            unidade.nome,
        ),
    )

    opcoes = []

    for unidade in unidades_ordenadas:
        rotulo = unidade.nome
        if unidade.sigla:
            rotulo = f"{unidade.nome} ({unidade.sigla})"

        opcoes.append(
            {
                "id": unidade.id,
                "codigo": unidade.codigo,
                "nome": unidade.nome,
                "rotulo": rotulo,
            }
        )

    return opcoes


def obter_unidade_por_correcao(request, numero_linha):
    unidade_id = request.POST.get(f"unidade_corrigida_{numero_linha}")

    if not unidade_id:
        return None

    return UnidadeMedida.objects.filter(id=unidade_id).first()


def buscar_item_existente(nome_item):
    nome_normalizado = normalizar_texto_comparacao(nome_item)

    if not nome_normalizado:
        return None

    item = Item.objects.filter(nome_item__iexact=nome_item).first()

    if item:
        return item

    for item in Item.objects.all():
        if normalizar_texto_comparacao(item.nome_item) == nome_normalizado:
            return item

    return None


def montar_linha_item(numero_linha, valores, atualizar_existentes):
    dados = dict(zip(CABECALHOS_ITENS, valores))

    nome_item = texto_celula(dados.get("NOME DO ITEM"))
    unidade_texto = texto_celula(dados.get("UNIDADE DE MEDIDA"))
    descricao = texto_celula(dados.get("DESCRIÇÃO"))
    ativo = texto_ativo(dados.get("ATIVO"))

    erros = []

    if not nome_item:
        erros.append("Nome do item é obrigatório.")

    if not unidade_texto:
        erros.append("Unidade de medida é obrigatória.")

    erro_unidade = False
    unidade_opcoes = []
    unidade = None

    if unidade_texto:
        unidade = buscar_unidade_medida_cadastrada(unidade_texto)

        if not unidade:
            erro_unidade = True
            unidade_opcoes = opcoes_unidade_para_correcao(unidade_texto)
            sugestoes = sugerir_unidades_medida(unidade_texto)

            if sugestoes:
                erros.append(
                    "Unidade de medida não encontrada no sistema. "
                    f"Selecione a unidade correta na coluna Correção. Sugestões: {', '.join(sugestoes)}."
                )
            else:
                erros.append(
                    "Unidade de medida não encontrada no sistema. "
                    "Selecione a unidade correta na coluna Correção ou cadastre a unidade antes de importar."
                )

    if not descricao:
        erros.append("Descrição é obrigatória.")

    item_existente = buscar_item_existente(nome_item)

    acao = "Criar"

    if item_existente:
        if atualizar_existentes:
            acao = "Atualizar"
        else:
            acao = "Ignorar"
            erros.append("Já existe um item cadastrado com este nome.")

    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "erro_unidade": erro_unidade,
        "unidade_opcoes": unidade_opcoes,
        "dados": {
            "nome_item": nome_item,
            "unidade_texto": unidade_texto,
            "unidade_medida": unidade.codigo if unidade else "",
            "unidade_corrigida_id": "",
            "unidade_corrigida_nome": "",
            "descricao": descricao,
            "ativo": ativo,
        },
    }


def ler_planilha_itens(arquivo, atualizar_existentes=True):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha_itens(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha os nomes das colunas."
        )

    linhas = []
    nomes_lidos = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_ITENS) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_item(
            numero_linha,
            valores,
            atualizar_existentes,
        )

        nome_normalizado = normalizar_texto_comparacao(linha["dados"]["nome_item"])

        if nome_normalizado:
            if nome_normalizado in nomes_lidos:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append("Nome do item duplicado dentro da própria planilha.")
            else:
                nomes_lidos.add(nome_normalizado)

        linhas.append(linha)

    return linhas


def criar_modelo_itens_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Itens"

    sheet.append(CABECALHOS_ITENS)
    sheet.append([
        "ARROZ TIPO 1",
        "Quilograma",
        "Arroz tipo 1, classe longo fino, pacote de 5kg, conforme especificação do edital.",
        "Sim",
    ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    larguras = {
        "A": 34,
        "B": 24,
        "C": 70,
        "D": 12,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 46

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_itens(request):
    workbook = criar_modelo_itens_workbook()
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_itens.xlsx"'

    return response


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_itens(request):
    linhas = []
    resumo = None

    if request.method == "POST":
        acao = request.POST.get("acao")

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]
                atualizar_existentes = form.cleaned_data["atualizar_existentes"]

                try:
                    linhas = ler_planilha_itens(
                        arquivo,
                        atualizar_existentes=atualizar_existentes,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [linha for linha in linhas if linha["valido"]]
                    erros = [linha for linha in linhas if not linha["valido"]]

                    request.session["importacao_itens_linhas"] = linhas
                    request.session["importacao_itens_atualizar"] = atualizar_existentes

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                        "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija as linhas inválidas antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )

        elif acao == "confirmar":
            linhas = request.session.get("importacao_itens_linhas", [])

            if not linhas:
                messages.error(request, "Nenhuma importação validada foi encontrada. Envie a planilha novamente.")
                return redirect("importacoes:itens")

            # Permite corrigir a unidade de medida diretamente na prévia.
            for linha in linhas:
                if linha.get("erro_unidade"):
                    unidade_corrigida = obter_unidade_por_correcao(request, linha["linha"])

                    if unidade_corrigida:
                        erros_restantes = [
                            erro for erro in linha.get("erros", [])
                            if "Unidade de medida não encontrada" not in erro
                        ]

                        linha["erros"] = erros_restantes
                        linha["erro_unidade"] = False
                        linha["dados"]["unidade_corrigida_id"] = unidade_corrigida.id
                        linha["dados"]["unidade_corrigida_nome"] = unidade_corrigida.nome
                        linha["dados"]["unidade_medida"] = unidade_corrigida.codigo
                        linha["dados"]["unidade_texto"] = unidade_corrigida.nome

                        if not erros_restantes:
                            linha["valido"] = True
                            if linha.get("acao") == "Erro":
                                linha["acao"] = "Atualizar" if buscar_item_existente(linha["dados"]["nome_item"]) else "Criar"

            linhas_invalidas = [linha for linha in linhas if not linha.get("valido")]

            if linhas_invalidas:
                request.session["importacao_itens_linhas"] = linhas
                messages.error(
                    request,
                    "Ainda existem linhas inválidas. Se o erro for unidade de medida, selecione a unidade correta na coluna Correção e confirme novamente.",
                )

                validos = [linha for linha in linhas if linha.get("valido")]
                erros = [linha for linha in linhas if not linha.get("valido")]
                resumo = {
                    "total": len(linhas),
                    "validos": len(validos),
                    "erros": len(erros),
                    "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                    "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                }

                form = ImportacaoPlanilhaForm()
                return render(
                    request,
                    "importacoes/itens.html",
                    {
                        "form": form,
                        "linhas": linhas,
                        "resumo": resumo,
                    },
                )

            criados = 0
            atualizados = 0

            for linha in linhas:
                dados = linha["dados"]

                unidade = None

                if dados.get("unidade_corrigida_id"):
                    unidade = UnidadeMedida.objects.filter(id=dados["unidade_corrigida_id"]).first()

                if not unidade and dados.get("unidade_medida"):
                    unidade = UnidadeMedida.objects.filter(codigo=dados["unidade_medida"]).first()

                if not unidade:
                    unidade = buscar_unidade_medida_cadastrada(dados["unidade_texto"])

                if not unidade:
                    messages.error(
                        request,
                        f"Unidade de medida não encontrada para o item {dados['nome_item']}. "
                        "Selecione a unidade correta na prévia ou corrija a planilha.",
                    )
                    return redirect("importacoes:itens")

                item = buscar_item_existente(dados["nome_item"])

                dados_item = {
                    "nome_item": dados["nome_item"],
                    "unidade_medida": unidade.codigo,
                    "descricao": dados["descricao"],
                    "ativo": dados["ativo"],
                }

                if item:
                    for campo, valor in dados_item.items():
                        setattr(item, campo, valor)
                    item.save()
                    atualizados += 1
                else:
                    Item.objects.create(**dados_item)
                    criados += 1

            request.session.pop("importacao_itens_linhas", None)
            request.session.pop("importacao_itens_atualizar", None)

            messages.success(
                request,
                f"Importação concluída. Itens criados: {criados}. "
                f"Itens atualizados: {atualizados}.",
            )
            return redirect("cadastros:itens")
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/itens.html",
        {
            "form": form,
            "linhas": linhas,
            "resumo": resumo,
        },
    )


CABECALHOS_QUANTITATIVO_PREGAO = [
    "ITEM",
    "UNIDADE DE MEDIDA",
    "QUANTIDADE DO MUNICÍPIO",
]


def validar_cabecalho_planilha_quantitativo_pregao(linha):
    encontrados = cabecalhos_normalizados(linha)

    return encontrados[:len(CABECALHOS_QUANTITATIVO_PREGAO)] == CABECALHOS_QUANTITATIVO_PREGAO


def decimal_planilha(valor):
    texto = texto_celula(valor)

    if not texto:
        return None

    texto = texto.replace(".", "").replace(",", ".") if "," in texto else texto

    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def opcoes_item_para_correcao(texto_item):
    itens = list(Item.objects.filter(ativo=True).order_by("nome_item"))

    sugestoes_normalizadas = set()
    texto_normalizado = normalizar_texto_comparacao(texto_item)

    mapa_normalizado = {
        normalizar_texto_comparacao(item.nome_item): item.nome_item
        for item in itens
    }

    sugestoes = get_close_matches(
        texto_normalizado,
        list(mapa_normalizado.keys()),
        n=8,
        cutoff=0.45,
    )

    for sugestao in sugestoes:
        sugestoes_normalizadas.add(sugestao)

    itens_ordenados = sorted(
        itens,
        key=lambda item: (
            0 if normalizar_texto_comparacao(item.nome_item) in sugestoes_normalizadas else 1,
            item.nome_item,
        ),
    )

    return [
        {
            "id": item.id,
            "nome": item.nome_item,
            "unidade": item.get_unidade_medida_display(),
        }
        for item in itens_ordenados
    ]


def obter_item_por_correcao(request, numero_linha):
    item_id = request.POST.get(f"item_corrigido_{numero_linha}")

    if not item_id:
        return None

    return Item.objects.filter(id=item_id).first()


def validar_item_unidade_quantitativo(item, unidade):
    if not item or not unidade:
        return False

    return item.unidade_medida == unidade.codigo


def montar_linha_quantitativo_pregao(numero_linha, valores, pregao, atualizar_existentes):
    dados = dict(zip(CABECALHOS_QUANTITATIVO_PREGAO, valores))

    item_texto = texto_celula(dados.get("ITEM"))
    unidade_texto = texto_celula(dados.get("UNIDADE DE MEDIDA"))
    quantidade = decimal_planilha(dados.get("QUANTIDADE DO MUNICÍPIO"))

    erros = []

    if not item_texto:
        erros.append("Item é obrigatório.")

    if not unidade_texto:
        erros.append("Unidade de medida é obrigatória.")

    if quantidade is None:
        erros.append("Quantidade do município é obrigatória e deve ser numérica.")
    elif quantidade <= 0:
        erros.append("Quantidade do município deve ser maior que zero.")

    item = None
    unidade = None
    erro_item = False
    erro_unidade = False
    item_opcoes = []
    unidade_opcoes = []

    if item_texto:
        item = buscar_item_existente(item_texto)

        if not item:
            erro_item = True
            item_opcoes = opcoes_item_para_correcao(item_texto)

            sugestoes_normalizadas = get_close_matches(
                normalizar_texto_comparacao(item_texto),
                [normalizar_texto_comparacao(i.nome_item) for i in Item.objects.filter(ativo=True)],
                n=5,
                cutoff=0.45,
            )

            if sugestoes_normalizadas:
                erros.append("Item não encontrado no sistema. Selecione o item correto na coluna Correção.")
            else:
                erros.append("Item não encontrado no sistema. Cadastre o item antes de importar ou selecione o item correto na coluna Correção.")

    if unidade_texto:
        unidade = buscar_unidade_medida_cadastrada(unidade_texto)

        if not unidade:
            erro_unidade = True
            unidade_opcoes = opcoes_unidade_para_correcao(unidade_texto)
            erros.append("Unidade de medida não encontrada no sistema. Selecione a unidade correta na coluna Correção.")

    if item and unidade and not validar_item_unidade_quantitativo(item, unidade):
        erro_unidade = True
        unidade_opcoes = opcoes_unidade_para_correcao(unidade_texto)
        erros.append(
            f"Unidade de medida não corresponde ao item cadastrado. "
            f"Unidade correta do item: {item.get_unidade_medida_display()}."
        )

    quantitativo_existente = None

    if pregao and item:
        quantitativo_existente = QuantitativoPregao.objects.filter(
            pregao=pregao,
            item=item,
        ).first()

    acao = "Criar"

    if quantitativo_existente:
        if atualizar_existentes:
            acao = "Atualizar"
        else:
            acao = "Ignorar"
            erros.append("Este item já possui quantitativo cadastrado para o pregão selecionado.")

    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "erro_item": erro_item,
        "erro_unidade": erro_unidade,
        "item_opcoes": item_opcoes,
        "unidade_opcoes": unidade_opcoes,
        "dados": {
            "item_texto": item_texto,
            "item_id": item.id if item else "",
            "item_corrigido_id": "",
            "item_corrigido_nome": "",
            "unidade_texto": unidade_texto,
            "unidade_medida": unidade.codigo if unidade else "",
            "unidade_corrigida_id": "",
            "unidade_corrigida_nome": "",
            "quantidade": str(quantidade) if quantidade is not None else "",
        },
    }


def ler_planilha_quantitativo_pregao(arquivo, pregao, atualizar_existentes=True):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha_quantitativo_pregao(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha os nomes das colunas."
        )

    linhas = []
    itens_lidos = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_QUANTITATIVO_PREGAO) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_quantitativo_pregao(
            numero_linha,
            valores,
            pregao,
            atualizar_existentes,
        )

        item_chave = linha["dados"].get("item_id") or normalizar_texto_comparacao(linha["dados"].get("item_texto"))

        if item_chave:
            if item_chave in itens_lidos:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append("Item duplicado dentro da própria planilha.")
            else:
                itens_lidos.add(item_chave)

        linhas.append(linha)

    return linhas


def criar_modelo_quantitativo_pregao_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quantitativo Pregão"

    sheet.append(CABECALHOS_QUANTITATIVO_PREGAO)
    sheet.append([
        "ARROZ TIPO 1",
        "Quilograma",
        "1000,000",
    ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    larguras = {
        "A": 44,
        "B": 24,
        "C": 24,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    sheet.row_dimensions[1].height = 32

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_quantitativo_pregao(request):
    workbook = criar_modelo_quantitativo_pregao_workbook()
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_quantitativo_pregao.xlsx"'

    return response


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_quantitativo_pregao(request):
    linhas = []
    resumo = None
    pregoes = Pregao.objects.filter(
        status=Pregao.STATUS_NAO_INICIADO,
    ).order_by("-ano", "-numero")
    pregao = None

    pregao_id = request.POST.get("pregao") or request.GET.get("pregao")

    if pregao_id:
        pregao = Pregao.objects.filter(id=pregao_id).first()

    if request.method == "POST":
        acao = request.POST.get("acao")

        if not pregao:
            messages.error(request, "Selecione um pregão válido.")
            return redirect("importacoes:quantitativo_pregao")

        if pregao.status != Pregao.STATUS_NAO_INICIADO:
            messages.error(
                request,
                "Não é permitido importar quantitativos em pregão já iniciado ou finalizado.",
            )
            return redirect("importacoes:quantitativo_pregao")

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]
                atualizar_existentes = form.cleaned_data["atualizar_existentes"]

                try:
                    linhas = ler_planilha_quantitativo_pregao(
                        arquivo,
                        pregao,
                        atualizar_existentes=atualizar_existentes,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [linha for linha in linhas if linha["valido"]]
                    erros = [linha for linha in linhas if not linha["valido"]]

                    request.session["importacao_quantitativo_pregao_linhas"] = linhas
                    request.session["importacao_quantitativo_pregao_pregao_id"] = pregao.id
                    request.session["importacao_quantitativo_pregao_atualizar"] = atualizar_existentes

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                        "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija as linhas inválidas antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )

        elif acao == "confirmar":
            linhas = request.session.get("importacao_quantitativo_pregao_linhas", [])
            pregao_sessao_id = request.session.get("importacao_quantitativo_pregao_pregao_id")

            if not linhas or not pregao_sessao_id:
                messages.error(request, "Nenhuma importação validada foi encontrada. Envie a planilha novamente.")
                return redirect("importacoes:quantitativo_pregao")

            pregao = Pregao.objects.filter(id=pregao_sessao_id).first()

            if not pregao:
                messages.error(request, "Pregão da importação não encontrado.")
                return redirect("importacoes:quantitativo_pregao")

            if pregao.status != Pregao.STATUS_NAO_INICIADO:
                messages.error(
                    request,
                    "Não é permitido importar quantitativos em pregão já iniciado ou finalizado.",
                )
                return redirect("importacoes:quantitativo_pregao")

            # Aplica correções de item e unidade feitas na tela de prévia.
            for linha in linhas:
                if linha.get("erro_item"):
                    item_corrigido = obter_item_por_correcao(request, linha["linha"])

                    if item_corrigido:
                        erros_restantes = [
                            erro for erro in linha.get("erros", [])
                            if "Item não encontrado" not in erro
                        ]

                        linha["erros"] = erros_restantes
                        linha["erro_item"] = False
                        linha["dados"]["item_id"] = item_corrigido.id
                        linha["dados"]["item_corrigido_id"] = item_corrigido.id
                        linha["dados"]["item_corrigido_nome"] = item_corrigido.nome_item
                        linha["dados"]["item_texto"] = item_corrigido.nome_item

                if linha.get("erro_unidade"):
                    unidade_corrigida = obter_unidade_por_correcao(request, linha["linha"])

                    if unidade_corrigida:
                        erros_restantes = [
                            erro for erro in linha.get("erros", [])
                            if "Unidade de medida não encontrada" not in erro
                            and "Unidade de medida não corresponde" not in erro
                        ]

                        linha["erros"] = erros_restantes
                        linha["erro_unidade"] = False
                        linha["dados"]["unidade_corrigida_id"] = unidade_corrigida.id
                        linha["dados"]["unidade_corrigida_nome"] = unidade_corrigida.nome
                        linha["dados"]["unidade_medida"] = unidade_corrigida.codigo
                        linha["dados"]["unidade_texto"] = unidade_corrigida.nome

                item = Item.objects.filter(id=linha["dados"].get("item_id")).first()
                unidade = None

                if linha["dados"].get("unidade_corrigida_id"):
                    unidade = UnidadeMedida.objects.filter(id=linha["dados"]["unidade_corrigida_id"]).first()

                if not unidade and linha["dados"].get("unidade_medida"):
                    unidade = UnidadeMedida.objects.filter(codigo=linha["dados"]["unidade_medida"]).first()

                if item and unidade:
                    erros_unidade_atuais = [
                        erro for erro in linha.get("erros", [])
                        if "Unidade de medida não corresponde" not in erro
                    ]

                    if validar_item_unidade_quantitativo(item, unidade):
                        linha["erros"] = erros_unidade_atuais
                        linha["erro_unidade"] = False
                    else:
                        linha["erros"] = erros_unidade_atuais + [
                            f"Unidade de medida não corresponde ao item cadastrado. Unidade correta do item: {item.get_unidade_medida_display()}."
                        ]
                        linha["erro_unidade"] = True

                if not linha.get("erros"):
                    linha["valido"] = True
                    linha["acao"] = "Atualizar" if QuantitativoPregao.objects.filter(
                        pregao=pregao,
                        item_id=linha["dados"].get("item_id"),
                    ).exists() else "Criar"
                else:
                    linha["valido"] = False
                    linha["acao"] = "Erro"

            # Verifica duplicidade após as correções.
            itens_finais = set()

            for linha in linhas:
                item_id = linha["dados"].get("item_id")

                if item_id:
                    if item_id in itens_finais:
                        linha["valido"] = False
                        linha["acao"] = "Erro"
                        linha["erros"].append("Item duplicado dentro da própria planilha após correção.")
                    else:
                        itens_finais.add(item_id)

            linhas_invalidas = [linha for linha in linhas if not linha.get("valido")]

            if linhas_invalidas:
                request.session["importacao_quantitativo_pregao_linhas"] = linhas

                messages.error(
                    request,
                    "Ainda existem linhas inválidas. Corrija a planilha ou selecione item/unidade correta na prévia e confirme novamente.",
                )

                validos = [linha for linha in linhas if linha.get("valido")]
                erros = [linha for linha in linhas if not linha.get("valido")]
                resumo = {
                    "total": len(linhas),
                    "validos": len(validos),
                    "erros": len(erros),
                    "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                    "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                }

                form = ImportacaoPlanilhaForm()
                return render(
                    request,
                    "importacoes/quantitativo_pregao.html",
                    {
                        "form": form,
                        "linhas": linhas,
                        "resumo": resumo,
                        "pregoes": pregoes,
                        "pregao": pregao,
                    },
                )

            criados = 0
            atualizados = 0

            for linha in linhas:
                dados = linha["dados"]
                item = Item.objects.get(id=dados["item_id"])
                quantidade = Decimal(str(dados["quantidade"]))

                quantitativo = QuantitativoPregao.objects.filter(
                    pregao=pregao,
                    item=item,
                ).first()

                if quantitativo:
                    quantitativo.quantidade = quantidade
                    quantitativo.full_clean()
                    quantitativo.save()
                    atualizados += 1
                else:
                    quantitativo = QuantitativoPregao(
                        pregao=pregao,
                        item=item,
                        quantidade=quantidade,
                    )
                    quantitativo.full_clean()
                    quantitativo.save()
                    criados += 1

            request.session.pop("importacao_quantitativo_pregao_linhas", None)
            request.session.pop("importacao_quantitativo_pregao_pregao_id", None)
            request.session.pop("importacao_quantitativo_pregao_atualizar", None)

            messages.success(
                request,
                f"Importação concluída para o Pregão {pregao.numero}/{pregao.ano}. "
                f"Itens criados: {criados}. Itens atualizados: {atualizados}.",
            )
            return redirect(f"/pregoes/quantitativo-pregao/?pregao={pregao.id}")
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/quantitativo_pregao.html",
        {
            "form": form,
            "linhas": linhas,
            "resumo": resumo,
            "pregoes": pregoes,
            "pregao": pregao,
        },
    )


CABECALHOS_QUANTITATIVO_ESCOLA = [
    "ESCOLA",
    "ITEM",
    "UNIDADE DE MEDIDA",
    "QUANTIDADE DA ESCOLA",
]


def validar_cabecalho_planilha_quantitativo_escola(linha):
    encontrados = cabecalhos_normalizados(linha)

    return encontrados[:len(CABECALHOS_QUANTITATIVO_ESCOLA)] == CABECALHOS_QUANTITATIVO_ESCOLA


def escolas_do_pregao(pregao):
    municipios_ids = pregao.municipios.values_list("id", flat=True)

    return (
        Escola.objects.filter(
            ativo=True,
            municipio_id__in=municipios_ids,
        )
        .select_related("municipio")
        .order_by("nome_escola")
    )


def buscar_escola_cadastrada_quantitativo(nome_escola, pregao):
    nome_escola = texto_celula(nome_escola)

    if not nome_escola or not pregao:
        return None

    escola = escolas_do_pregao(pregao).filter(nome_escola__iexact=nome_escola).first()

    if escola:
        return escola

    nome_normalizado = normalizar_texto_comparacao(nome_escola)

    for escola in escolas_do_pregao(pregao):
        if normalizar_texto_comparacao(escola.nome_escola) == nome_normalizado:
            return escola

    return None


def opcoes_escola_para_correcao(nome_escola, pregao):
    escolas = list(escolas_do_pregao(pregao))

    texto_normalizado = normalizar_texto_comparacao(nome_escola)

    mapa_normalizado = {
        normalizar_texto_comparacao(escola.nome_escola): escola.nome_escola
        for escola in escolas
    }

    sugestoes = get_close_matches(
        texto_normalizado,
        list(mapa_normalizado.keys()),
        n=8,
        cutoff=0.45,
    )

    sugestoes_normalizadas = set(sugestoes)

    escolas_ordenadas = sorted(
        escolas,
        key=lambda escola: (
            0 if normalizar_texto_comparacao(escola.nome_escola) in sugestoes_normalizadas else 1,
            escola.nome_escola,
        ),
    )

    return [
        {
            "id": escola.id,
            "nome": escola.nome_escola,
            "municipio": escola.municipio.nome if escola.municipio else "",
            "uf": escola.municipio.uf if escola.municipio else "",
        }
        for escola in escolas_ordenadas
    ]


def obter_escola_por_correcao(request, numero_linha):
    escola_id = request.POST.get(f"escola_corrigida_{numero_linha}")

    if not escola_id:
        return None

    return Escola.objects.filter(id=escola_id).first()


def item_existe_no_quantitativo_pregao(pregao, item):
    if not pregao or not item:
        return False

    return QuantitativoPregao.objects.filter(
        pregao=pregao,
        item=item,
    ).exists()


def obter_quantitativo_pregao(pregao, item):
    if not pregao or not item:
        return None

    return QuantitativoPregao.objects.filter(
        pregao=pregao,
        item=item,
    ).first()


def saldo_disponivel_quantitativo_escola(pregao, escola, item):
    quantitativo_pregao = obter_quantitativo_pregao(pregao, item)

    if not quantitativo_pregao:
        return Decimal("0")

    total_pregao = quantitativo_pregao.quantidade

    total_outras_escolas = (
        QuantitativoEscola.objects.filter(
            pregao=pregao,
            item=item,
        )
        .exclude(escola=escola)
        .aggregate(total=Sum("quantidade"))
        .get("total")
    )

    if total_outras_escolas is None:
        total_outras_escolas = Decimal("0")

    return total_pregao - total_outras_escolas


def validar_escola_pertence_pregao(pregao, escola):
    if not pregao or not escola:
        return False

    return pregao.municipios.filter(id=escola.municipio_id).exists()


def montar_linha_quantitativo_escola(numero_linha, valores, pregao, atualizar_existentes):
    dados = dict(zip(CABECALHOS_QUANTITATIVO_ESCOLA, valores))

    escola_texto = texto_celula(dados.get("ESCOLA"))
    item_texto = texto_celula(dados.get("ITEM"))
    unidade_texto = texto_celula(dados.get("UNIDADE DE MEDIDA"))
    quantidade = decimal_planilha(dados.get("QUANTIDADE DA ESCOLA"))

    erros = []

    if not escola_texto:
        erros.append("Escola é obrigatória.")

    if not item_texto:
        erros.append("Item é obrigatório.")

    if not unidade_texto:
        erros.append("Unidade de medida é obrigatória.")

    if quantidade is None:
        erros.append("Quantidade da escola é obrigatória e deve ser numérica.")
    elif quantidade <= 0:
        erros.append("Quantidade da escola deve ser maior que zero.")

    escola = None
    item = None
    unidade = None

    erro_escola = False
    erro_item = False
    erro_unidade = False

    escola_opcoes = []
    item_opcoes = []
    unidade_opcoes = []

    if escola_texto:
        escola = buscar_escola_cadastrada_quantitativo(escola_texto, pregao)

        if not escola:
            erro_escola = True
            escola_opcoes = opcoes_escola_para_correcao(escola_texto, pregao)
            erros.append(
                "Escola não encontrada entre as escolas dos municípios vinculados ao pregão. "
                "Selecione a escola correta na coluna Correção."
            )
        elif not validar_escola_pertence_pregao(pregao, escola):
            erro_escola = True
            escola_opcoes = opcoes_escola_para_correcao(escola_texto, pregao)
            erros.append("A escola informada não pertence aos municípios vinculados ao pregão selecionado.")

    if item_texto:
        item = buscar_item_existente(item_texto)

        if not item:
            erro_item = True
            item_opcoes = opcoes_item_para_correcao(item_texto)
            erros.append("Item não encontrado no sistema. Selecione o item correto na coluna Correção.")
        elif not item_existe_no_quantitativo_pregao(pregao, item):
            erro_item = True
            item_opcoes = opcoes_item_para_correcao(item_texto)
            erros.append("Este item não existe no quantitativo do pregão selecionado.")

    if unidade_texto:
        unidade = buscar_unidade_medida_cadastrada(unidade_texto)

        if not unidade:
            erro_unidade = True
            unidade_opcoes = opcoes_unidade_para_correcao(unidade_texto)
            erros.append("Unidade de medida não encontrada no sistema. Selecione a unidade correta na coluna Correção.")

    if item and unidade and not validar_item_unidade_quantitativo(item, unidade):
        erro_unidade = True
        unidade_opcoes = opcoes_unidade_para_correcao(unidade_texto)
        erros.append(
            f"Unidade de medida não corresponde ao item cadastrado. "
            f"Unidade correta do item: {item.get_unidade_medida_display()}."
        )

    if pregao and escola and item and quantidade and quantidade > 0:
        saldo = saldo_disponivel_quantitativo_escola(pregao, escola, item)

        if quantidade > saldo:
            erros.append(
                f"Quantidade da escola ultrapassa o saldo disponível para este item. "
                f"Saldo máximo permitido: {saldo}."
            )

    quantitativo_existente = None

    if pregao and escola and item:
        quantitativo_existente = QuantitativoEscola.objects.filter(
            pregao=pregao,
            escola=escola,
            item=item,
        ).first()

    acao = "Criar"

    if quantitativo_existente:
        if atualizar_existentes:
            acao = "Atualizar"
        else:
            acao = "Ignorar"
            erros.append("Esta escola já possui quantitativo cadastrado para este item no pregão selecionado.")

    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "erro_escola": erro_escola,
        "erro_item": erro_item,
        "erro_unidade": erro_unidade,
        "escola_opcoes": escola_opcoes,
        "item_opcoes": item_opcoes,
        "unidade_opcoes": unidade_opcoes,
        "dados": {
            "escola_texto": escola_texto,
            "escola_id": escola.id if escola else "",
            "escola_corrigida_id": "",
            "escola_corrigida_nome": "",
            "item_texto": item_texto,
            "item_id": item.id if item else "",
            "item_corrigido_id": "",
            "item_corrigido_nome": "",
            "unidade_texto": unidade_texto,
            "unidade_medida": unidade.codigo if unidade else "",
            "unidade_corrigida_id": "",
            "unidade_corrigida_nome": "",
            "quantidade": str(quantidade) if quantidade is not None else "",
        },
    }


def ler_planilha_quantitativo_escola(arquivo, pregao, atualizar_existentes=True):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha_quantitativo_escola(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha os nomes das colunas."
        )

    linhas = []
    chaves_lidas = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_QUANTITATIVO_ESCOLA) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_quantitativo_escola(
            numero_linha,
            valores,
            pregao,
            atualizar_existentes,
        )

        escola_chave = linha["dados"].get("escola_id") or normalizar_texto_comparacao(linha["dados"].get("escola_texto"))
        item_chave = linha["dados"].get("item_id") or normalizar_texto_comparacao(linha["dados"].get("item_texto"))
        chave = f"{escola_chave}|{item_chave}"

        if escola_chave and item_chave:
            if chave in chaves_lidas:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append("Combinação escola/item duplicada dentro da própria planilha.")
            else:
                chaves_lidas.add(chave)

        linhas.append(linha)

    return linhas


def criar_modelo_quantitativo_escola_workbook():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quantitativo Escola"

    sheet.append(CABECALHOS_QUANTITATIVO_ESCOLA)
    sheet.append([
        "EE EXEMPLO",
        "ARROZ TIPO 1",
        "Quilograma",
        "100,000",
    ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2, max_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    larguras = {
        "A": 44,
        "B": 44,
        "C": 24,
        "D": 24,
    }

    for coluna, largura in larguras.items():
        sheet.column_dimensions[coluna].width = largura

    sheet.row_dimensions[1].height = 32

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_quantitativo_escola(request):
    workbook = criar_modelo_quantitativo_escola_workbook()
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modelo_importacao_quantitativo_escola.xlsx"'

    return response


def recalcular_linha_quantitativo_escola(linha, pregao):
    erros = linha.get("erros", [])

    escola = Escola.objects.filter(id=linha["dados"].get("escola_id")).first()
    item = Item.objects.filter(id=linha["dados"].get("item_id")).first()
    unidade = None

    if linha["dados"].get("unidade_corrigida_id"):
        unidade = UnidadeMedida.objects.filter(id=linha["dados"]["unidade_corrigida_id"]).first()

    if not unidade and linha["dados"].get("unidade_medida"):
        unidade = UnidadeMedida.objects.filter(codigo=linha["dados"]["unidade_medida"]).first()

    erros_filtrados = []
    for erro in erros:
        if "Escola não encontrada" in erro:
            continue
        if "escola informada não pertence" in erro:
            continue
        if "Item não encontrado" in erro:
            continue
        if "não existe no quantitativo do pregão" in erro:
            continue
        if "Unidade de medida não encontrada" in erro:
            continue
        if "Unidade de medida não corresponde" in erro:
            continue
        if "Quantidade da escola ultrapassa" in erro:
            continue
        if "já possui quantitativo cadastrado" in erro:
            continue
        erros_filtrados.append(erro)

    if escola:
        if not validar_escola_pertence_pregao(pregao, escola):
            linha["erro_escola"] = True
            linha["escola_opcoes"] = opcoes_escola_para_correcao(linha["dados"].get("escola_texto"), pregao)
            erros_filtrados.append("A escola selecionada não pertence aos municípios vinculados ao pregão selecionado.")
        else:
            linha["erro_escola"] = False
    else:
        linha["erro_escola"] = True
        erros_filtrados.append("Escola não encontrada entre as escolas dos municípios vinculados ao pregão.")

    if item:
        if not item_existe_no_quantitativo_pregao(pregao, item):
            linha["erro_item"] = True
            linha["item_opcoes"] = opcoes_item_para_correcao(linha["dados"].get("item_texto"))
            erros_filtrados.append("Este item não existe no quantitativo do pregão selecionado.")
        else:
            linha["erro_item"] = False
    else:
        linha["erro_item"] = True
        erros_filtrados.append("Item não encontrado no sistema.")

    if unidade:
        if item and not validar_item_unidade_quantitativo(item, unidade):
            linha["erro_unidade"] = True
            linha["unidade_opcoes"] = opcoes_unidade_para_correcao(linha["dados"].get("unidade_texto"))
            erros_filtrados.append(
                f"Unidade de medida não corresponde ao item cadastrado. "
                f"Unidade correta do item: {item.get_unidade_medida_display()}."
            )
        else:
            linha["erro_unidade"] = False
    else:
        linha["erro_unidade"] = True
        erros_filtrados.append("Unidade de medida não encontrada no sistema.")

    quantidade = decimal_planilha(linha["dados"].get("quantidade"))

    if quantidade is None:
        erros_filtrados.append("Quantidade da escola é obrigatória e deve ser numérica.")
    elif quantidade <= 0:
        erros_filtrados.append("Quantidade da escola deve ser maior que zero.")
    elif pregao and escola and item and item_existe_no_quantitativo_pregao(pregao, item):
        saldo = saldo_disponivel_quantitativo_escola(pregao, escola, item)

        if quantidade > saldo:
            erros_filtrados.append(
                f"Quantidade da escola ultrapassa o saldo disponível para este item. "
                f"Saldo máximo permitido: {saldo}."
            )

    quantitativo_existente = None

    if pregao and escola and item:
        quantitativo_existente = QuantitativoEscola.objects.filter(
            pregao=pregao,
            escola=escola,
            item=item,
        ).first()

    if quantitativo_existente:
        linha["acao"] = "Atualizar"
    else:
        linha["acao"] = "Criar"

    linha["erros"] = erros_filtrados
    linha["valido"] = not erros_filtrados
    if not linha["valido"]:
        linha["acao"] = "Erro"

    return linha


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_quantitativo_escola(request):
    linhas = []
    resumo = None
    pregoes = Pregao.objects.filter(
        status=Pregao.STATUS_NAO_INICIADO,
    ).order_by("-ano", "-numero")
    pregao = None

    pregao_id = request.POST.get("pregao") or request.GET.get("pregao")

    if pregao_id:
        pregao = Pregao.objects.filter(id=pregao_id).first()

    if request.method == "POST":
        acao = request.POST.get("acao")

        if not pregao:
            messages.error(request, "Selecione um pregão válido.")
            return redirect("importacoes:quantitativo_escola")

        if pregao.status != Pregao.STATUS_NAO_INICIADO:
            messages.error(
                request,
                "Não é permitido importar quantitativos de escola em pregão já iniciado ou finalizado.",
            )
            return redirect("importacoes:quantitativo_escola")

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]
                atualizar_existentes = form.cleaned_data["atualizar_existentes"]

                try:
                    linhas = ler_planilha_quantitativo_escola(
                        arquivo,
                        pregao,
                        atualizar_existentes=atualizar_existentes,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [linha for linha in linhas if linha["valido"]]
                    erros = [linha for linha in linhas if not linha["valido"]]

                    request.session["importacao_quantitativo_escola_linhas"] = linhas
                    request.session["importacao_quantitativo_escola_pregao_id"] = pregao.id
                    request.session["importacao_quantitativo_escola_atualizar"] = atualizar_existentes

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                        "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija as linhas inválidas antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )

        elif acao == "confirmar":
            linhas = request.session.get("importacao_quantitativo_escola_linhas", [])
            pregao_sessao_id = request.session.get("importacao_quantitativo_escola_pregao_id")

            if not linhas or not pregao_sessao_id:
                messages.error(request, "Nenhuma importação validada foi encontrada. Envie a planilha novamente.")
                return redirect("importacoes:quantitativo_escola")

            pregao = Pregao.objects.filter(id=pregao_sessao_id).first()

            if not pregao:
                messages.error(request, "Pregão da importação não encontrado.")
                return redirect("importacoes:quantitativo_escola")

            if pregao.status != Pregao.STATUS_NAO_INICIADO:
                messages.error(
                    request,
                    "Não é permitido importar quantitativos de escola em pregão já iniciado ou finalizado.",
                )
                return redirect("importacoes:quantitativo_escola")

            for linha in linhas:
                if linha.get("erro_escola"):
                    escola_corrigida = obter_escola_por_correcao(request, linha["linha"])

                    if escola_corrigida:
                        linha["dados"]["escola_id"] = escola_corrigida.id
                        linha["dados"]["escola_corrigida_id"] = escola_corrigida.id
                        linha["dados"]["escola_corrigida_nome"] = escola_corrigida.nome_escola
                        linha["dados"]["escola_texto"] = escola_corrigida.nome_escola

                if linha.get("erro_item"):
                    item_corrigido = obter_item_por_correcao(request, linha["linha"])

                    if item_corrigido:
                        linha["dados"]["item_id"] = item_corrigido.id
                        linha["dados"]["item_corrigido_id"] = item_corrigido.id
                        linha["dados"]["item_corrigido_nome"] = item_corrigido.nome_item
                        linha["dados"]["item_texto"] = item_corrigido.nome_item

                if linha.get("erro_unidade"):
                    unidade_corrigida = obter_unidade_por_correcao(request, linha["linha"])

                    if unidade_corrigida:
                        linha["dados"]["unidade_corrigida_id"] = unidade_corrigida.id
                        linha["dados"]["unidade_corrigida_nome"] = unidade_corrigida.nome
                        linha["dados"]["unidade_medida"] = unidade_corrigida.codigo
                        linha["dados"]["unidade_texto"] = unidade_corrigida.nome

                linha = recalcular_linha_quantitativo_escola(linha, pregao)

            # Verifica duplicidade depois das correções.
            chaves_finais = set()

            for linha in linhas:
                escola_id = linha["dados"].get("escola_id")
                item_id = linha["dados"].get("item_id")

                if escola_id and item_id:
                    chave = f"{escola_id}|{item_id}"

                    if chave in chaves_finais:
                        linha["valido"] = False
                        linha["acao"] = "Erro"
                        linha["erros"].append("Combinação escola/item duplicada dentro da própria planilha após correção.")
                    else:
                        chaves_finais.add(chave)

            linhas_invalidas = [linha for linha in linhas if not linha.get("valido")]

            if linhas_invalidas:
                request.session["importacao_quantitativo_escola_linhas"] = linhas

                messages.error(
                    request,
                    "Ainda existem linhas inválidas. Corrija a planilha ou selecione escola/item/unidade correta na prévia e confirme novamente.",
                )

                validos = [linha for linha in linhas if linha.get("valido")]
                erros = [linha for linha in linhas if not linha.get("valido")]
                resumo = {
                    "total": len(linhas),
                    "validos": len(validos),
                    "erros": len(erros),
                    "criar": len([linha for linha in validos if linha["acao"] == "Criar"]),
                    "atualizar": len([linha for linha in validos if linha["acao"] == "Atualizar"]),
                }

                form = ImportacaoPlanilhaForm()
                return render(
                    request,
                    "importacoes/quantitativo_escola.html",
                    {
                        "form": form,
                        "linhas": linhas,
                        "resumo": resumo,
                        "pregoes": pregoes,
                        "pregao": pregao,
                    },
                )

            criados = 0
            atualizados = 0

            for linha in linhas:
                dados = linha["dados"]
                escola = Escola.objects.get(id=dados["escola_id"])
                item = Item.objects.get(id=dados["item_id"])
                quantidade = Decimal(str(dados["quantidade"]))

                quantitativo = QuantitativoEscola.objects.filter(
                    pregao=pregao,
                    escola=escola,
                    item=item,
                ).first()

                if quantitativo:
                    quantitativo.quantidade = quantidade
                    quantitativo.full_clean()
                    quantitativo.save()
                    atualizados += 1
                else:
                    quantitativo = QuantitativoEscola(
                        pregao=pregao,
                        escola=escola,
                        item=item,
                        quantidade=quantidade,
                    )
                    quantitativo.full_clean()
                    quantitativo.save()
                    criados += 1

            request.session.pop("importacao_quantitativo_escola_linhas", None)
            request.session.pop("importacao_quantitativo_escola_pregao_id", None)
            request.session.pop("importacao_quantitativo_escola_atualizar", None)

            messages.success(
                request,
                f"Importação concluída para o Pregão {pregao.numero}/{pregao.ano}. "
                f"Quantitativos criados: {criados}. Quantitativos atualizados: {atualizados}.",
            )
            return redirect("pregoes:quantitativo_escola")
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/quantitativo_escola.html",
        {
            "form": form,
            "linhas": linhas,
            "resumo": resumo,
            "pregoes": pregoes,
            "pregao": pregao,
        },
    )


# ============================================================
# ETAPA 75 — EXPORTAÇÃO / BACKUP DE DADOS
# ============================================================

def valor_sim_nao(valor):
    return "Sim" if valor else "Não"


def texto_data(valor):
    if not valor:
        return ""

    try:
        return valor.strftime("%d/%m/%Y")
    except Exception:
        return str(valor)


def texto_data_hora(valor):
    if not valor:
        return ""

    try:
        return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M")
    except Exception:
        return str(valor)


def formatar_numero_decimal(valor):
    if valor is None:
        return ""

    return str(valor).replace(".", ",")


def aplicar_estilo_backup(sheet):
    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    if sheet.max_row >= 1:
        for cell in sheet[1]:
            cell.fill = fill_header
            cell.font = font_header
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.freeze_panes = "A2"

    for coluna in range(1, sheet.max_column + 1):
        letra = get_column_letter(coluna)
        largura = 12

        for cell in sheet[letra]:
            valor = "" if cell.value is None else str(cell.value)
            largura = max(largura, min(len(valor) + 3, 55))

        sheet.column_dimensions[letra].width = largura

    sheet.row_dimensions[1].height = 32


def adicionar_aba(workbook, titulo, cabecalhos, linhas):
    sheet = workbook.create_sheet(title=titulo)
    sheet.append(cabecalhos)

    for linha in linhas:
        sheet.append(linha)

    aplicar_estilo_backup(sheet)

    return sheet


def criar_backup_completo_workbook(usuario=None):
    workbook = Workbook()

    # Remove a aba padrão para criar tudo em ordem.
    aba_padrao = workbook.active
    workbook.remove(aba_padrao)

    agora = timezone.localtime()

    resumo = workbook.create_sheet(title="Resumo")
    resumo.append(["Informação", "Valor"])
    resumo.append(["Gerado em", agora.strftime("%d/%m/%Y %H:%M")])
    nome_usuario = ""
    if usuario:
        nome_usuario = usuario.get_full_name() or usuario.username

    resumo.append(["Gerado por", nome_usuario])
    resumo.append(["Municípios", Municipio.objects.count()])
    resumo.append(["Unidades de Medida", UnidadeMedida.objects.count()])
    resumo.append(["Fornecedores", Fornecedor.objects.count()])
    resumo.append(["Escolas", Escola.objects.count()])
    resumo.append(["Itens", Item.objects.count()])
    resumo.append(["Pregões", Pregao.objects.count()])
    resumo.append(["Quantitativos por Pregão", QuantitativoPregao.objects.count()])
    resumo.append(["Quantitativos por Escola", QuantitativoEscola.objects.count()])
    aplicar_estilo_backup(resumo)

    adicionar_aba(
        workbook,
        "Municípios",
        ["ID", "Nome", "UF", "Ativo", "Criado em", "Atualizado em"],
        [
            [
                municipio.id,
                municipio.nome,
                municipio.uf,
                valor_sim_nao(municipio.ativo),
                texto_data_hora(municipio.criado_em),
                texto_data_hora(municipio.atualizado_em),
            ]
            for municipio in Municipio.objects.all().order_by("nome", "uf")
        ],
    )

    adicionar_aba(
        workbook,
        "Unidades de Medida",
        ["ID", "Nome", "Sigla", "Código", "Ativo", "Criado em", "Atualizado em"],
        [
            [
                unidade.id,
                unidade.nome,
                unidade.sigla,
                unidade.codigo,
                valor_sim_nao(unidade.ativo),
                texto_data_hora(unidade.criado_em),
                texto_data_hora(unidade.atualizado_em),
            ]
            for unidade in UnidadeMedida.objects.all().order_by("nome")
        ],
    )

    adicionar_aba(
        workbook,
        "Fornecedores",
        [
            "ID",
            "Razão Social",
            "CNPJ",
            "Endereço",
            "Telefone",
            "Representante Legal",
            "RG Representante",
            "Órgão Expedidor",
            "CPF Representante",
            "E-mail",
            "Banco",
            "Agência",
            "Conta Corrente",
            "PIX",
            "Ativo",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                fornecedor.id,
                fornecedor.razao_social,
                fornecedor.cnpj,
                fornecedor.endereco,
                fornecedor.telefone,
                fornecedor.representante_legal,
                fornecedor.rg_representante,
                fornecedor.orgao_expedidor_representante,
                fornecedor.cpf_representante,
                fornecedor.email,
                fornecedor.nome_banco,
                fornecedor.agencia,
                fornecedor.conta_corrente,
                fornecedor.pix,
                valor_sim_nao(fornecedor.ativo),
                texto_data_hora(fornecedor.criado_em),
                texto_data_hora(fornecedor.atualizado_em),
            ]
            for fornecedor in Fornecedor.objects.all().order_by("razao_social")
        ],
    )

    adicionar_aba(
        workbook,
        "Escolas",
        [
            "ID",
            "Nome da Escola",
            "CNPJ",
            "Endereço",
            "Número",
            "Bairro",
            "Município",
            "UF",
            "Presidente CDCE",
            "RG Presidente",
            "CPF Presidente",
            "Ativo",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                escola.id,
                escola.nome_escola,
                escola.cnpj,
                escola.endereco,
                escola.numero,
                escola.bairro,
                escola.municipio.nome if escola.municipio else "",
                escola.municipio.uf if escola.municipio else "",
                escola.presidente_cdce,
                escola.rg_presidente,
                escola.cpf_presidente,
                valor_sim_nao(escola.ativo),
                texto_data_hora(escola.criado_em),
                texto_data_hora(escola.atualizado_em),
            ]
            for escola in Escola.objects.select_related("municipio").all().order_by("nome_escola")
        ],
    )

    adicionar_aba(
        workbook,
        "Itens",
        [
            "ID",
            "Nome do Item",
            "Unidade Código",
            "Unidade",
            "Descrição",
            "Ativo",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                item.id,
                item.nome_item,
                item.unidade_medida,
                item.get_unidade_medida_display(),
                item.descricao,
                valor_sim_nao(item.ativo),
                texto_data_hora(item.criado_em),
                texto_data_hora(item.atualizado_em),
            ]
            for item in Item.objects.all().order_by("nome_item")
        ],
    )

    adicionar_aba(
        workbook,
        "Pregões",
        [
            "ID",
            "Número",
            "Ano",
            "Municípios",
            "Fornecedores",
            "Pregoeiro",
            "CPF Pregoeiro",
            "Local",
            "Data",
            "Status",
            "Criado por",
            "Finalizado em",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                pregao.id,
                pregao.numero,
                pregao.ano,
                ", ".join(str(municipio) for municipio in pregao.municipios.all().order_by("nome")),
                ", ".join(fornecedor.razao_social for fornecedor in pregao.fornecedores.all().order_by("razao_social")),
                pregao.nome_pregoeiro,
                pregao.cpf_pregoeiro,
                pregao.local_pregao,
                texto_data(pregao.data_pregao),
                pregao.get_status_display(),
                pregao.criado_por.get_full_name() or pregao.criado_por.username if pregao.criado_por else "",
                texto_data_hora(pregao.finalizado_em),
                texto_data_hora(pregao.criado_em),
                texto_data_hora(pregao.atualizado_em),
            ]
            for pregao in Pregao.objects.all().prefetch_related("municipios", "fornecedores").select_related("criado_por").order_by("-ano", "-numero")
        ],
    )

    adicionar_aba(
        workbook,
        "Quantitativo Pregão",
        [
            "ID",
            "Pregão",
            "Ano",
            "Item",
            "Unidade",
            "Quantidade",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                quantitativo.id,
                quantitativo.pregao.numero,
                quantitativo.pregao.ano,
                quantitativo.item.nome_item,
                quantitativo.item.get_unidade_medida_display(),
                formatar_numero_decimal(quantitativo.quantidade),
                texto_data_hora(quantitativo.criado_em),
                texto_data_hora(quantitativo.atualizado_em),
            ]
            for quantitativo in QuantitativoPregao.objects.select_related("pregao", "item").all().order_by("-pregao__ano", "-pregao__numero", "item__nome_item")
        ],
    )

    adicionar_aba(
        workbook,
        "Quantitativo Escola",
        [
            "ID",
            "Pregão",
            "Ano",
            "Escola",
            "Município",
            "Item",
            "Unidade",
            "Quantidade",
            "Criado em",
            "Atualizado em",
        ],
        [
            [
                quantitativo.id,
                quantitativo.pregao.numero,
                quantitativo.pregao.ano,
                quantitativo.escola.nome_escola,
                str(quantitativo.escola.municipio) if quantitativo.escola and quantitativo.escola.municipio else "",
                quantitativo.item.nome_item,
                quantitativo.item.get_unidade_medida_display(),
                formatar_numero_decimal(quantitativo.quantidade),
                texto_data_hora(quantitativo.criado_em),
                texto_data_hora(quantitativo.atualizado_em),
            ]
            for quantitativo in QuantitativoEscola.objects.select_related("pregao", "escola", "escola__municipio", "item").all().order_by("-pregao__ano", "-pregao__numero", "escola__nome_escola", "item__nome_item")
        ],
    )

    workbook.active = 0

    return workbook


# ============================================================
# IMPORTAÇÃO DE MÉDIA DE PREÇOS
# ============================================================

CABECALHOS_MEDIA_PRECOS = [
    "ITEM",
    "MÉDIA DE PREÇO",
]


def validar_cabecalho_planilha_media_precos(linha):
    encontrados = cabecalhos_normalizados(linha)
    return encontrados[:len(CABECALHOS_MEDIA_PRECOS)] == CABECALHOS_MEDIA_PRECOS


def garantir_itens_pregao_para_importacao_media(pregao):
    """
    Garante que os PregaoItem existam antes da importação das médias,
    usando a mesma regra da tela manual de Média de Preços:
    só cria os itens se houver quantitativo por escola para o certame.
    """
    if PregaoItem.objects.filter(pregao=pregao).exists():
        return True

    existe_quantitativo_escola = QuantitativoEscola.objects.filter(
        pregao=pregao
    ).exists()

    if not existe_quantitativo_escola:
        return False

    quantitativos_pregao = (
        QuantitativoPregao.objects.filter(pregao=pregao)
        .select_related("item")
        .order_by("item__nome_item")
    )

    ordem = 1

    for qp in quantitativos_pregao:
        PregaoItem.objects.get_or_create(
            pregao=pregao,
            item=qp.item,
            defaults={
                "ordem": ordem,
                "quantidade_total": qp.quantidade,
            },
        )
        ordem += 1

    return PregaoItem.objects.filter(pregao=pregao).exists()


def opcoes_item_certame_para_correcao(texto_item, pregao):
    """
    Retorna somente itens que pertencem ao certame selecionado.
    Mantém os itens mais parecidos com o texto informado no topo da lista.
    """
    itens_pregao = (
        PregaoItem.objects.filter(pregao=pregao)
        .select_related("item")
        .order_by("ordem", "item__nome_item")
    )

    itens = [registro.item for registro in itens_pregao]

    texto_normalizado = normalizar_texto_comparacao(texto_item)

    mapa_normalizado = {
        normalizar_texto_comparacao(item.nome_item): item
        for item in itens
    }

    sugestoes = set(
        get_close_matches(
            texto_normalizado,
            list(mapa_normalizado.keys()),
            n=8,
            cutoff=0.45,
        )
    )

    itens_ordenados = sorted(
        itens,
        key=lambda item: (
            0 if normalizar_texto_comparacao(item.nome_item) in sugestoes else 1,
            item.nome_item,
        ),
    )

    return [
        {
            "id": item.id,
            "nome": item.nome_item,
            "unidade": item.get_unidade_medida_display(),
        }
        for item in itens_ordenados
    ]


def montar_linha_media_precos(numero_linha, valores, pregao):
    dados = dict(zip(CABECALHOS_MEDIA_PRECOS, valores))

    item_texto = texto_celula(dados.get("ITEM"))
    media_preco = decimal_planilha(dados.get("MÉDIA DE PREÇO"))

    erros = []
    erro_item = False
    item_opcoes = []
    item = None
    item_pregao = None

    if not item_texto:
        erros.append("Item é obrigatório.")

    if media_preco is None:
        erros.append("Média de preço é obrigatória e deve ser numérica.")
    elif media_preco <= 0:
        erros.append("Média de preço deve ser maior que zero.")

    if item_texto:
        item = buscar_item_existente(item_texto)

        if not item:
            erro_item = True
            item_opcoes = opcoes_item_certame_para_correcao(item_texto, pregao)
            erros.append(
                "Item não encontrado no sistema. Selecione o item correto na coluna Correção."
            )
        else:
            item_pregao = PregaoItem.objects.filter(
                pregao=pregao,
                item=item,
            ).first()

            if not item_pregao:
                erro_item = True
                item_opcoes = opcoes_item_certame_para_correcao(item_texto, pregao)
                erros.append(
                    "O item existe no cadastro, mas não pertence ao certame selecionado. "
                    "Selecione um item correto do certame na coluna Correção ou marque Não importar."
                )

    acao = "Atualizar" if item_pregao else "Erro"
    valido = not erros

    return {
        "linha": numero_linha,
        "valido": valido,
        "acao": acao if valido else "Erro",
        "erros": erros,
        "erro_item": erro_item,
        "item_opcoes": item_opcoes,
        "dados": {
            "item_texto": item_texto,
            "item_id": item.id if item else "",
            "item_pregao_id": item_pregao.id if item_pregao else "",
            "item_corrigido_id": "",
            "item_corrigido_nome": "",
            "media_preco": str(media_preco) if media_preco is not None else "",
        },
    }


def ler_planilha_media_precos(arquivo, pregao):
    workbook = load_workbook(arquivo, data_only=True)
    sheet = workbook.active

    primeira_linha = [cell.value for cell in sheet[1]]

    if not validar_cabecalho_planilha_media_precos(primeira_linha):
        raise ValueError(
            "O cabeçalho da planilha não corresponde ao modelo esperado. "
            "Baixe o modelo novamente e mantenha as colunas ITEM e MÉDIA DE PREÇO."
        )

    linhas = []
    itens_lidos = set()

    for numero_linha in range(2, sheet.max_row + 1):
        valores = [
            sheet.cell(row=numero_linha, column=coluna).value
            for coluna in range(1, len(CABECALHOS_MEDIA_PRECOS) + 1)
        ]

        if all(texto_celula(valor) == "" for valor in valores):
            continue

        linha = montar_linha_media_precos(
            numero_linha,
            valores,
            pregao,
        )

        item_chave = (
            linha["dados"].get("item_id")
            or normalizar_texto_comparacao(linha["dados"].get("item_texto"))
        )

        if item_chave:
            if item_chave in itens_lidos:
                linha["valido"] = False
                linha["acao"] = "Erro"
                linha["erros"].append(
                    "Item duplicado dentro da própria planilha."
                )
            else:
                itens_lidos.add(item_chave)

        linhas.append(linha)

    return linhas


def criar_modelo_media_precos_workbook(pregao=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Média de Preços"

    sheet.append(CABECALHOS_MEDIA_PRECOS)

    if pregao:
        garantir_itens_pregao_para_importacao_media(pregao)

        itens_pregao = (
            PregaoItem.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("ordem", "item__nome_item")
        )

        for item_pregao in itens_pregao:
            sheet.append([
                item_pregao.item.nome_item,
                (
                    float(item_pregao.media_preco)
                    if item_pregao.media_preco is not None
                    else None
                ),
            ])
    else:
        sheet.append([
            "ARROZ BRANCO",
            7.50,
        ])

    sheet.freeze_panes = "A2"

    fill_header = PatternFill("solid", fgColor="0F2F57")
    font_header = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in sheet[1]:
        cell.fill = fill_header
        cell.font = font_header
        cell.border = border
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 22
    sheet.row_dimensions[1].height = 32

    for cell in sheet["B"][1:]:
        cell.number_format = 'R$ #,##0.00'

    return workbook


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_modelo_media_precos(request):
    pregao = None
    pregao_id = request.GET.get("pregao")

    if pregao_id:
        pregao = Pregao.objects.filter(id=pregao_id).first()

    workbook = criar_modelo_media_precos_workbook(pregao)
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    nome_arquivo = "modelo_importacao_media_precos.xlsx"

    if pregao:
        nome_arquivo = (
            f"media_precos_{pregao.numero}_{pregao.ano}.xlsx"
            .replace("/", "-")
            .replace("\\", "-")
        )

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{nome_arquivo}"'
    )

    return response


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def importar_media_precos(request):
    linhas = []
    resumo = None

    pregoes = Pregao.objects.filter(
        status__in=[
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]
    ).order_by("-ano", "-numero")

    pregao = None
    pregao_id = request.POST.get("pregao") or request.GET.get("pregao")

    if pregao_id:
        pregao = Pregao.objects.filter(id=pregao_id).first()

    if request.method == "POST":
        acao = request.POST.get("acao")

        if not pregao:
            messages.error(request, "Selecione um certame válido.")
            return redirect("importacoes:media_precos")

        if pregao.status not in [
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]:
            messages.error(
                request,
                "Só é possível importar médias em certames não iniciados ou em andamento.",
            )
            return redirect("importacoes:media_precos")

        if not garantir_itens_pregao_para_importacao_media(pregao):
            messages.error(
                request,
                "Não foi possível localizar itens para este certame. "
                "Cadastre primeiro o quantitativo do certame e o quantitativo por escola.",
            )
            return redirect(
                f"{request.path}?pregao={pregao.id}"
            )

        if acao == "validar":
            form = ImportacaoPlanilhaForm(request.POST, request.FILES)

            if form.is_valid():
                arquivo = form.cleaned_data["arquivo"]

                try:
                    linhas = ler_planilha_media_precos(
                        arquivo,
                        pregao,
                    )
                except Exception as erro:
                    messages.error(request, str(erro))
                    linhas = []

                if linhas:
                    validos = [
                        linha for linha in linhas if linha["valido"]
                    ]
                    erros = [
                        linha for linha in linhas if not linha["valido"]
                    ]

                    request.session["importacao_media_precos_linhas"] = linhas
                    request.session["importacao_media_precos_pregao_id"] = pregao.id

                    resumo = {
                        "total": len(linhas),
                        "validos": len(validos),
                        "erros": len(erros),
                        "atualizar": len(validos),
                        "ignorados": 0,
                    }

                    if erros:
                        messages.warning(
                            request,
                            "A planilha possui erros. Corrija os itens indicados antes de confirmar a importação.",
                        )
                    else:
                        messages.success(
                            request,
                            "Planilha validada com sucesso. Confira a prévia e confirme a importação.",
                        )

        elif acao == "confirmar":
            linhas = request.session.get(
                "importacao_media_precos_linhas",
                [],
            )
            pregao_sessao_id = request.session.get(
                "importacao_media_precos_pregao_id"
            )

            if not linhas or not pregao_sessao_id:
                messages.error(
                    request,
                    "Nenhuma importação validada foi encontrada. Envie a planilha novamente.",
                )
                return redirect("importacoes:media_precos")

            pregao = Pregao.objects.filter(
                id=pregao_sessao_id
            ).first()

            if not pregao:
                messages.error(
                    request,
                    "Certame da importação não encontrado.",
                )
                return redirect("importacoes:media_precos")

            # Marca as linhas que o usuário escolheu não importar.
            for linha in linhas:
                linha["ignorar"] = (
                    request.POST.get(f"ignorar_linha_{linha['linha']}") == "on"
                )

            # Revalida cada linha a partir do estado atual da prévia.
            # Isso evita manter erros antigos depois que o usuário corrige um item.
            for linha in linhas:
                if linha.get("ignorar"):
                    linha["valido"] = True
                    linha["acao"] = "Ignorar"
                    continue

                dados = linha.get("dados", {})
                erros_novos = []

                # Mantém somente erros de média, pois os erros de item serão
                # recalculados abaixo com base na correção escolhida.
                for erro in linha.get("erros", []):
                    if (
                        "Média de preço" in erro
                        or "média de preço" in erro
                        or "MÉDIA DE PREÇO" in erro
                    ):
                        erros_novos.append(erro)

                item_corrigido = obter_item_por_correcao(
                    request,
                    linha["linha"],
                )

                item = None
                item_pregao = None

                if item_corrigido:
                    item = item_corrigido
                else:
                    item_id = dados.get("item_id")
                    if item_id:
                        item = Item.objects.filter(id=item_id, ativo=True).first()

                if not item:
                    erros_novos.append(
                        "Item não encontrado no sistema. "
                        "Selecione o item correto na coluna Correção ou marque Não importar."
                    )
                    linha["erro_item"] = True
                    linha["item_opcoes"] = opcoes_item_certame_para_correcao(
                        dados.get("item_texto", ""),
                        pregao,
                    )
                    dados["item_pregao_id"] = ""
                    dados["item_corrigido_id"] = ""
                    dados["item_corrigido_nome"] = ""
                else:
                    item_pregao = PregaoItem.objects.filter(
                        pregao=pregao,
                        item=item,
                    ).first()

                    if not item_pregao:
                        erros_novos.append(
                            "O item selecionado não pertence ao certame. "
                            "Selecione outro item na coluna Correção ou marque Não importar."
                        )
                        linha["erro_item"] = True
                        linha["item_opcoes"] = opcoes_item_para_correcao(
                            dados.get("item_texto", "")
                        )
                        dados["item_pregao_id"] = ""
                    else:
                        linha["erro_item"] = False
                        dados["item_id"] = item.id
                        dados["item_pregao_id"] = item_pregao.id

                        if item_corrigido:
                            dados["item_corrigido_id"] = item.id
                            dados["item_corrigido_nome"] = item.nome_item
                            dados["item_texto"] = item.nome_item

                linha["dados"] = dados
                linha["erros"] = erros_novos
                linha["valido"] = not erros_novos
                linha["acao"] = "Atualizar" if linha["valido"] else "Erro"

            # Verifica duplicidade somente após aplicar todas as correções e ignorados.
            # Assim, duplicidades antigas não ficam "presas" depois de uma correção.
            itens_confirmacao = {}

            for linha in linhas:
                if linha.get("ignorar") or not linha.get("valido"):
                    continue

                item_pregao_id = linha["dados"].get("item_pregao_id")

                if not item_pregao_id:
                    linha["valido"] = False
                    linha["acao"] = "Erro"
                    linha["erros"].append(
                        "Não foi possível identificar o item do certame."
                    )
                    continue

                if item_pregao_id in itens_confirmacao:
                    primeira_linha = itens_confirmacao[item_pregao_id]
                    linha["valido"] = False
                    linha["acao"] = "Erro"
                    linha["erros"].append(
                        f"O mesmo item também foi informado na linha {primeira_linha}. "
                        "Marque uma das linhas como Não importar."
                    )
                else:
                    itens_confirmacao[item_pregao_id] = linha["linha"]

            linhas_invalidas = [
                linha
                for linha in linhas
                if not linha.get("ignorar") and not linha.get("valido")
            ]

            if linhas_invalidas:
                request.session["importacao_media_precos_linhas"] = linhas

                validos = [
                    linha
                    for linha in linhas
                    if not linha.get("ignorar") and linha.get("valido")
                ]
                ignorados = [
                    linha for linha in linhas if linha.get("ignorar")
                ]

                resumo = {
                    "total": len(linhas),
                    "validos": len(validos),
                    "erros": len(linhas_invalidas),
                    "atualizar": len(validos),
                    "ignorados": len(ignorados),
                }

                linhas_com_erro = ", ".join(
                    str(linha["linha"]) for linha in linhas_invalidas
                )

                messages.error(
                    request,
                    f"Ainda existem linhas inválidas: {linhas_com_erro}. "
                    "Corrija o item indicado ou marque Não importar e confirme novamente.",
                )

                form = ImportacaoPlanilhaForm()
                return render(
                    request,
                    "importacoes/media_precos.html",
                    {
                        "form": form,
                        "pregoes": pregoes,
                        "pregao": pregao,
                        "linhas": linhas,
                        "resumo": resumo,
                    },
                )

            atualizados = 0

            ignorados = 0

            for linha in linhas:
                if linha.get("ignorar"):
                    ignorados += 1
                    continue

                dados = linha["dados"]

                item_pregao = PregaoItem.objects.filter(
                    id=dados["item_pregao_id"],
                    pregao=pregao,
                ).first()

                if not item_pregao:
                    continue

                item_pregao.media_preco = Decimal(
                    dados["media_preco"]
                )
                item_pregao.save(
                    update_fields=["media_preco"]
                )
                atualizados += 1

            request.session.pop(
                "importacao_media_precos_linhas",
                None,
            )
            request.session.pop(
                "importacao_media_precos_pregao_id",
                None,
            )

            messages.success(
                request,
                f"Importação concluída. Médias atualizadas: {atualizados}. "
                f"Itens ignorados: {ignorados}.",
            )

            return redirect(
                f"/pregoes/media-precos/?pregao={pregao.id}"
            )
    else:
        form = ImportacaoPlanilhaForm()

    if request.method == "POST" and "form" not in locals():
        form = ImportacaoPlanilhaForm()

    return render(
        request,
        "importacoes/media_precos.html",
        {
            "form": form,
            "pregoes": pregoes,
            "pregao": pregao,
            "linhas": linhas,
            "resumo": resumo,
        },
    )


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def exportacoes(request):
    return render(request, "importacoes/exportacoes.html")


@login_required
@user_passes_test(usuario_pode_importar, login_url="dashboard")
def baixar_backup_completo(request):
    workbook = criar_backup_completo_workbook(usuario=request.user)
    arquivo = BytesIO()
    workbook.save(arquivo)
    arquivo.seek(0)

    data_arquivo = timezone.localtime().strftime("%Y-%m-%d_%H-%M")
    nome_arquivo = f"backup_sistema_pregoes_{data_arquivo}.xlsx"

    response = HttpResponse(
        arquivo.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'

    return response
