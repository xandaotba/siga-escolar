from decimal import Decimal, InvalidOperation
import base64
import binascii
import hashlib
import hmac
import json
import re
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core import signing
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db import IntegrityError, transaction
from django.db.models import Sum
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_GET, require_POST

from django.http import HttpResponse, JsonResponse
from apps.cadastros.models import Escola, Item, validar_cnpj
from .forms import OrdenadorDespesasForm, PregaoForm

from io import BytesIO
from copy import deepcopy
from pathlib import Path

from django.conf import settings
from docx import Document

def paginar_queryset(request, queryset, por_pagina=25):
    paginator = Paginator(queryset, por_pagina)
    numero_pagina = request.GET.get("page")
    return paginator.get_page(numero_pagina)


from .models import (
    Pregao,
    PregaoItem,
    PregaoFornecedor,
    QuantitativoEscola,
    QuantitativoPregao,
    PropostaInicialItem,
    PropostaDigitalLink,
    PropostaDigital,
)



def _substituir_texto_celula_preservando_rotulo(cell, valor, rotulo=None):
    """
    Preenche uma célula do modelo mantendo o rótulo em negrito quando existir.
    """
    valor = str(valor or "").strip() or "-"

    paragraph = cell.paragraphs[0]

    if rotulo:
        # Mantém o primeiro run (rótulo) do modelo e substitui apenas o conteúdo.
        if not paragraph.runs:
            run_rotulo = paragraph.add_run(rotulo)
            run_rotulo.bold = True
            paragraph.add_run(f" {valor}")
            return

        paragraph.runs[0].text = rotulo
        paragraph.runs[0].bold = True

        if len(paragraph.runs) >= 2:
            paragraph.runs[1].text = f" {valor}"
            for run in paragraph.runs[2:]:
                run.text = ""
        else:
            paragraph.add_run(f" {valor}")
    else:
        if paragraph.runs:
            paragraph.runs[0].text = valor
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(valor)


def _preencher_tabela_fornecedor(table, fornecedor):
    _substituir_texto_celula_preservando_rotulo(
        table.cell(0, 1),
        fornecedor.razao_social,
    )
    _substituir_texto_celula_preservando_rotulo(
        table.cell(1, 1),
        fornecedor.cnpj,
    )
    _substituir_texto_celula_preservando_rotulo(
        table.cell(2, 1),
        fornecedor.endereco,
    )
    _substituir_texto_celula_preservando_rotulo(
        table.cell(3, 1),
        fornecedor.representante_legal,
        "NOME:",
    )
    _substituir_texto_celula_preservando_rotulo(
        table.cell(4, 1),
        fornecedor.cpf_representante,
        "CPF:",
    )
    identidade = fornecedor.rg_representante or ""
    orgao = fornecedor.orgao_expedidor_representante or ""

    if identidade and orgao:
        identidade_completa = f"{identidade} - {orgao}"
    else:
        identidade_completa = identidade or orgao

    _substituir_texto_celula_preservando_rotulo(
        table.cell(5, 1),
        identidade_completa,
        "IDENTIDADE:",
    )
    _substituir_texto_celula_preservando_rotulo(
        table.cell(6, 1),
        fornecedor.telefone,
    )


def relacao_fornecedores_word(request, pregao_id):
    pregao = get_object_or_404(Pregao, id=pregao_id)

    # Este documento é exclusivo dos Pregões Presenciais.
    if not pregao.eh_pregao_presencial:
        messages.error(
            request,
            "A Relação de Fornecedores está disponível somente para Pregões Presenciais.",
        )
        return redirect("pregoes:pregoes")

    fornecedores_pregao = list(
        PregaoFornecedor.objects.filter(
            pregao=pregao,
            ativo_no_pregao=True,
        )
        .select_related("fornecedor")
        .order_by("ordem_inicial", "fornecedor__razao_social")
    )

    if not fornecedores_pregao:
        messages.warning(
            request,
            "Não há fornecedores cadastrados para este pregão.",
        )
        return redirect("pregoes:pregoes")

    caminho_modelo = (
        Path(settings.BASE_DIR)
        / "templates"
        / "documentos"
        / "modelo_relacao_fornecedores_pregao.docx"
    )

    if not caminho_modelo.exists():
        messages.error(
            request,
            "O modelo Word da Relação de Fornecedores não foi encontrado no servidor.",
        )
        return redirect("pregoes:pregoes")

    document = Document(caminho_modelo)

    if not document.tables:
        messages.error(
            request,
            "O modelo Word da Relação de Fornecedores não possui a tabela esperada.",
        )
        return redirect("pregoes:pregoes")

    tabela_modelo = document.tables[0]

    # Guarda uma cópia da tabela original antes do preenchimento.
    tabela_xml_modelo = deepcopy(tabela_modelo._tbl)

    # Primeiro fornecedor usa a própria tabela existente no modelo.
    _preencher_tabela_fornecedor(
        tabela_modelo,
        fornecedores_pregao[0].fornecedor,
    )

    # Os demais fornecedores recebem uma cópia idêntica da tabela do modelo.
    elemento_anterior = tabela_modelo._tbl

    for vinculo in fornecedores_pregao[1:]:
        # Espaço entre as tabelas.
        paragrafo_espaco = document.add_paragraph()
        elemento_anterior.addnext(paragrafo_espaco._p)
        elemento_anterior = paragrafo_espaco._p

        nova_tabela_xml = deepcopy(tabela_xml_modelo)
        elemento_anterior.addnext(nova_tabela_xml)
        elemento_anterior = nova_tabela_xml

        # Localiza a tabela recém-inserida para preencher seus dados.
        nova_tabela = document.tables[-1]
        _preencher_tabela_fornecedor(
            nova_tabela,
            vinculo.fornecedor,
        )

    arquivo = BytesIO()
    document.save(arquivo)
    arquivo.seek(0)

    nome_arquivo = (
        f"Relacao_Fornecedores_Pregao_{pregao.numero}_{pregao.ano}.docx"
        .replace("/", "-")
        .replace("\\", "-")
    )

    response = HttpResponse(
        arquivo.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
    )
    response["Content-Disposition"] = f'attachment; filename="{nome_arquivo}"'

    return response



def pregoes(request):
    if request.method == "POST":
        form = PregaoForm(request.POST)

        if form.is_valid():
            form.save(user=request.user)
            messages.success(request, "Certame cadastrado com sucesso.")
            return redirect("pregoes:pregoes")

        messages.error(request, "Verifique os dados informados no formulário.")
    else:
        form = PregaoForm()

    pregoes_queryset = (
        Pregao.objects.all()
        .prefetch_related("municipios", "fornecedores")
        .order_by("-ano", "-numero")
    )
    pregoes_pagina = paginar_queryset(request, pregoes_queryset)

    return render(
        request,
        "pregoes/pregoes.html",
        {
            "form": form,
            "pregoes": pregoes_pagina,
            "modo_edicao": False,
            "pregao_edicao": None,
        },
    )


def editar_pregao(request, pregao_id):
    pregao = get_object_or_404(Pregao, id=pregao_id)

    if pregao.status != Pregao.STATUS_NAO_INICIADO:
        messages.error(
            request,
            "Este certame não pode ser editado porque já está em andamento ou finalizado.",
        )
        return redirect("pregoes:pregoes")

    if request.method == "POST":
        form = PregaoForm(request.POST, instance=pregao)

        if form.is_valid():
            form.save(user=request.user)
            messages.success(request, "Certame atualizado com sucesso.")
            return redirect("pregoes:pregoes")

        messages.error(request, "Verifique os dados informados no formulário.")
    else:
        form = PregaoForm(instance=pregao)

    pregoes_queryset = (
        Pregao.objects.all()
        .prefetch_related("municipios", "fornecedores")
        .order_by("-ano", "-numero")
    )
    pregoes_pagina = paginar_queryset(request, pregoes_queryset)

    return render(
        request,
        "pregoes/pregoes.html",
        {
            "form": form,
            "pregoes": pregoes_pagina,
            "modo_edicao": True,
            "pregao_edicao": pregao,
        },
    )



def editar_ordenador_despesas(request, pregao_id):
    """
    Permite cadastrar/alterar somente Nome, RG e CPF do ordenador de despesas.

    Esta edição permanece disponível para certames em qualquer status porque
    não altera dados da disputa, quantitativos, fornecedores, municípios ou
    resultados do certame.
    """
    pregao = get_object_or_404(Pregao, id=pregao_id)

    if request.method == "POST":
        form = OrdenadorDespesasForm(request.POST, instance=pregao)

        if form.is_valid():
            form.save()
            messages.success(
                request,
                "Dados do ordenador de despesas atualizados com sucesso.",
            )
            return redirect("pregoes:pregoes")

        messages.error(
            request,
            "Verifique os dados informados para o ordenador de despesas.",
        )
    else:
        form = OrdenadorDespesasForm(instance=pregao)

    return render(
        request,
        "pregoes/ordenador_despesas.html",
        {
            "form": form,
            "pregao": pregao,
        },
    )


def _converter_quantidade_digitada(valor, valor_atual=None):
    """
    Converte quantidades digitadas aceitando padrão brasileiro e o formato
    técnico usado pelo próprio SIGA.

    Exemplos:
    1442 -> 1442
    1.442 -> 1442
    1.442,500 -> 1442.500
    1442,500 -> 1442.500
    1,442 -> 1.442
    1442.500 -> 1442.500

    Se o valor já veio preenchido pelo próprio sistema e não foi alterado,
    preserva exatamente o Decimal existente.
    """
    texto = str(valor or "").strip()
    texto = texto.replace("\xa0", "").replace(" ", "")

    if not texto:
        return None

    if valor_atual is not None:
        try:
            atual = Decimal(str(valor_atual))
            if texto in {str(atual), format(atual, "f")}:
                return atual
        except (InvalidOperation, ValueError, TypeError):
            pass

    if not re.fullmatch(r"[+-]?[0-9][0-9.,]*", texto):
        raise InvalidOperation

    sinal = ""
    if texto[0] in "+-":
        sinal = texto[0]
        texto = texto[1:]

    if "," in texto:
        # Padrão brasileiro: ponto = milhar, vírgula = decimal.
        if texto.count(",") > 1:
            raise InvalidOperation

        parte_inteira, parte_decimal = texto.split(",", 1)

        if "." in parte_decimal:
            raise InvalidOperation

        parte_inteira = parte_inteira.replace(".", "") or "0"
        normalizado = parte_inteira

        if parte_decimal:
            normalizado += "." + parte_decimal

    elif "." in texto:
        partes = texto.split(".")

        if len(partes) > 2:
            # Ex.: 1.234.567
            if not partes[0] or not all(
                grupo.isdigit() and len(grupo) == 3
                for grupo in partes[1:]
            ):
                raise InvalidOperation
            normalizado = "".join(partes)

        else:
            esquerda, direita = partes

            if not esquerda or not direita:
                raise InvalidOperation

            # Ex.: 1.442 / 12.345 / 123.456 = separador de milhar.
            # Ex.: 1442.500 = representação decimal técnica.
            if len(esquerda) <= 3 and len(direita) == 3:
                normalizado = esquerda + direita
            else:
                normalizado = texto

    else:
        normalizado = texto

    return Decimal(sinal + normalizado)


def quantitativo_pregao(request):
    pregoes_lista = (
        Pregao.objects.filter(
            status=Pregao.STATUS_NAO_INICIADO
        )
        .prefetch_related("municipios")
        .order_by("-ano", "-numero")
    )

    pregao_id = request.GET.get("pregao") or request.POST.get("pregao")
    pregao = None
    itens = Item.objects.filter(ativo=True).order_by("nome_item")
    quantitativos_salvos = {}

    if pregao_id:
        pregao = get_object_or_404(Pregao, id=pregao_id)

        quantitativos_salvos = {
            q.item_id: q.quantidade
            for q in QuantitativoPregao.objects.filter(pregao=pregao)
        }

    if request.method == "POST":
        if not pregao:
            messages.error(request, "Selecione um pregão.")
            return redirect("pregoes:quantitativo_pregao")

        if pregao.status != Pregao.STATUS_NAO_INICIADO:
            messages.error(
                request,
                "Não é permitido alterar quantitativos de pregão já iniciado ou finalizado.",
            )
            return redirect("pregoes:quantitativo_pregao")

        quantidades_digitadas = {}
        erros_validacao = []

        for item in itens:
            quantidade_texto = request.POST.get(
                f"quantidade_{item.id}",
                "",
            ).strip()

            if not quantidade_texto:
                continue

            try:
                quantidade = _converter_quantidade_digitada(
                    quantidade_texto,
                    valor_atual=quantitativos_salvos.get(item.id),
                )
            except (InvalidOperation, ValueError, TypeError):
                erros_validacao.append(
                    f"A quantidade informada para o item '{item.nome_item}' é inválida."
                )
                continue

            if quantidade is not None and quantidade > 0:
                quantidades_digitadas[item.id] = quantidade

        if erros_validacao:
            for erro in erros_validacao:
                messages.error(request, erro)

            quantitativos_salvos = {
                **quantitativos_salvos,
                **quantidades_digitadas,
            }

            return render(
                request,
                "pregoes/quantitativo_pregao.html",
                {
                    "pregoes": pregoes_lista,
                    "pregao": pregao,
                    "itens": itens,
                    "quantitativos_salvos": quantitativos_salvos,
                },
            )

        with transaction.atomic():
            QuantitativoPregao.objects.filter(pregao=pregao).delete()

            for item_id, quantidade in quantidades_digitadas.items():
                QuantitativoPregao.objects.create(
                    pregao=pregao,
                    item_id=item_id,
                    quantidade=quantidade,
                )

        total_salvos = len(quantidades_digitadas)

        messages.success(
            request,
            f"Quantitativos salvos com sucesso. Itens gravados: {total_salvos}.",
        )
        return redirect(f"{request.path}?pregao={pregao.id}")

    return render(
        request,
        "pregoes/quantitativo_pregao.html",
        {
            "pregoes": pregoes_lista,
            "pregao": pregao,
            "itens": itens,
            "quantitativos_salvos": quantitativos_salvos,
        },
    )


def quantitativo_escola(request):
    pregoes_lista = (
        Pregao.objects.filter(
            status=Pregao.STATUS_NAO_INICIADO
        )
        .prefetch_related("municipios")
        .order_by("-ano", "-numero")
    )

    pregao_id = request.GET.get("pregao") or request.POST.get("pregao")
    escola_id = request.GET.get("escola") or request.POST.get("escola")

    pregao = None
    escola = None
    escolas = Escola.objects.none()
    itens_pregao = []
    quantitativos_salvos = {}
    resumo_itens = []

    if pregao_id:
        pregao = get_object_or_404(Pregao, id=pregao_id)

        municipios_ids = pregao.municipios.values_list("id", flat=True)

        escolas = (
            Escola.objects.filter(
                ativo=True,
                municipio_id__in=municipios_ids,
            )
            .select_related("municipio")
            .order_by("nome_escola")
        )

    if escola_id:
        escola = get_object_or_404(Escola, id=escola_id)

    if pregao and escola:
        if not pregao.municipios.filter(id=escola.municipio_id).exists():
            messages.error(
                request,
                "A escola selecionada não pertence aos municípios vinculados a este pregão.",
            )
            return redirect("pregoes:quantitativo_escola")

    if pregao:
        itens_pregao = (
            QuantitativoPregao.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("item__nome_item")
        )

    if pregao and escola:
        quantitativos_salvos = {
            q.item_id: q.quantidade
            for q in QuantitativoEscola.objects.filter(
                pregao=pregao,
                escola=escola,
            )
        }

        for qp in itens_pregao:
            item = qp.item
            quantidade_total_pregao = qp.quantidade

            quantidade_escola_atual = quantitativos_salvos.get(
                item.id,
                Decimal("0"),
            )

            distribuido_todas_escolas = (
                QuantitativoEscola.objects.filter(
                    pregao=pregao,
                    item=item,
                )
                .aggregate(total=Sum("quantidade"))
                .get("total")
            )

            if distribuido_todas_escolas is None:
                distribuido_todas_escolas = Decimal("0")

            saldo_restante_real = quantidade_total_pregao - distribuido_todas_escolas

            resumo_itens.append(
                {
                    "item": item,
                    "quantidade_total_pregao": quantidade_total_pregao,
                    "quantidade_escola_atual": quantidade_escola_atual,
                    "distribuido_todas_escolas": distribuido_todas_escolas,
                    "saldo_restante_real": saldo_restante_real,
                }
            )

    if request.method == "POST":
        if not pregao or not escola:
            messages.error(request, "Selecione o pregão e a escola.")
            return redirect("pregoes:quantitativo_escola")

        if pregao.status != Pregao.STATUS_NAO_INICIADO:
            messages.error(
                request,
                "Não é permitido alterar quantitativos de escola em pregão já iniciado ou finalizado.",
            )
            return redirect("pregoes:quantitativo_escola")

        if not pregao.municipios.filter(id=escola.municipio_id).exists():
            messages.error(
                request,
                "A escola selecionada não pertence aos municípios vinculados a este pregão.",
            )
            return redirect("pregoes:quantitativo_escola")

        quantidades_digitadas = {}
        erros_validacao = []

        for qp in itens_pregao:
            item = qp.item
            quantidade_texto = request.POST.get(
                f"quantidade_{item.id}",
                "",
            ).strip()

            if not quantidade_texto:
                continue

            try:
                quantidade_nova = _converter_quantidade_digitada(
                    quantidade_texto,
                    valor_atual=quantitativos_salvos.get(item.id),
                )
            except (InvalidOperation, ValueError, TypeError):
                erros_validacao.append(
                    f"A quantidade informada para o item '{item.nome_item}' é inválida."
                )
                continue

            if quantidade_nova is None or quantidade_nova <= 0:
                continue

            quantidade_total_pregao = qp.quantidade

            soma_outras_escolas = (
                QuantitativoEscola.objects.filter(
                    pregao=pregao,
                    item=item,
                )
                .exclude(escola=escola)
                .aggregate(total=Sum("quantidade"))
                .get("total")
            )

            if soma_outras_escolas is None:
                soma_outras_escolas = Decimal("0")

            soma_final = soma_outras_escolas + quantidade_nova

            if soma_final > quantidade_total_pregao:
                saldo_permitido_para_esta_escola = (
                    quantidade_total_pregao - soma_outras_escolas
                )

                erros_validacao.append(
                    f"O item '{item.nome_item}' ultrapassou o total do pregão. "
                    f"Total do pregão: {quantidade_total_pregao}. "
                    f"Quantidade máxima permitida para esta escola: "
                    f"{saldo_permitido_para_esta_escola}."
                )
            else:
                quantidades_digitadas[item.id] = quantidade_nova

        if erros_validacao:
            for erro in erros_validacao:
                messages.error(request, erro)

            resumo_itens = []

            for qp in itens_pregao:
                item = qp.item
                quantidade_total_pregao = qp.quantidade

                quantidade_escola_atual = quantidades_digitadas.get(
                    item.id,
                    quantitativos_salvos.get(item.id, Decimal("0")),
                )

                distribuido_todas_escolas = (
                    QuantitativoEscola.objects.filter(
                        pregao=pregao,
                        item=item,
                    )
                    .exclude(escola=escola)
                    .aggregate(total=Sum("quantidade"))
                    .get("total")
                )

                if distribuido_todas_escolas is None:
                    distribuido_todas_escolas = Decimal("0")

                distribuido_todas_escolas = (
                    distribuido_todas_escolas + quantidade_escola_atual
                )
                saldo_restante_real = (
                    quantidade_total_pregao - distribuido_todas_escolas
                )

                resumo_itens.append(
                    {
                        "item": item,
                        "quantidade_total_pregao": quantidade_total_pregao,
                        "quantidade_escola_atual": quantidade_escola_atual,
                        "distribuido_todas_escolas": distribuido_todas_escolas,
                        "saldo_restante_real": saldo_restante_real,
                    }
                )

            return render(
                request,
                "pregoes/quantitativo_escola.html",
                {
                    "pregoes": pregoes_lista,
                    "escolas": escolas,
                    "pregao": pregao,
                    "escola": escola,
                    "itens_pregao": itens_pregao,
                    "quantitativos_salvos": quantidades_digitadas,
                    "resumo_itens": resumo_itens,
                },
            )

        with transaction.atomic():
            QuantitativoEscola.objects.filter(
                pregao=pregao,
                escola=escola,
            ).delete()

            for item_id, quantidade in quantidades_digitadas.items():
                QuantitativoEscola.objects.create(
                    pregao=pregao,
                    escola=escola,
                    item_id=item_id,
                    quantidade=quantidade,
                )

        total_salvos = len(quantidades_digitadas)

        messages.success(
            request,
            f"Quantitativos da escola salvos com sucesso. Itens gravados: {total_salvos}.",
        )

        return redirect(
            f"{request.path}?pregao={pregao.id}&escola={escola.id}"
        )

    return render(
        request,
        "pregoes/quantitativo_escola.html",
        {
            "pregoes": pregoes_lista,
            "escolas": escolas,
            "pregao": pregao,
            "escola": escola,
            "itens_pregao": itens_pregao,
            "quantitativos_salvos": quantitativos_salvos,
            "resumo_itens": resumo_itens,
        },
    )


def usuario_eh_consulta_escola(request):
    perfil = getattr(request.user, "perfil_acesso", None)

    return bool(
        request.user.is_authenticated
        and perfil
        and perfil.perfil == "consulta_escola"
    )


def obter_escola_consulta_usuario(request):
    perfil = getattr(request.user, "perfil_acesso", None)

    if not perfil:
        return None

    return perfil.escola


def relatorio_distribuicao(request):
    consulta_escola = usuario_eh_consulta_escola(request)
    escola_consulta = obter_escola_consulta_usuario(request) if consulta_escola else None

    if consulta_escola:
        if escola_consulta:
            pregoes_ids = (
                QuantitativoEscola.objects.filter(
                    escola=escola_consulta,
                    quantidade__gt=0,
                )
                .values_list("pregao_id", flat=True)
                .distinct()
            )

            pregoes_lista = Pregao.objects.filter(
                id__in=pregoes_ids,
            ).order_by("-ano", "-numero")
        else:
            pregoes_lista = Pregao.objects.none()

            messages.warning(
                request,
                "Seu usuário está com o perfil Consulta/Escola, mas não possui escola vinculada. "
                "Solicite ao administrador que vincule seu usuário a uma escola.",
            )
    else:
        pregoes_lista = Pregao.objects.all().order_by("-ano", "-numero")

    pregao_id = request.GET.get("pregao")
    pregao = None
    relatorio = []

    total_itens = 0
    itens_fechados = 0
    itens_com_saldo = 0
    itens_excedidos = 0

    if pregao_id:
        # Segurança: usuário Consulta/Escola só pode abrir relatório de pregões
        # que possuam quantitativo vinculado à sua própria escola.
        if consulta_escola:
            if not escola_consulta:
                messages.error(
                    request,
                    "Seu usuário não possui escola vinculada para consultar relatórios.",
                )
                return redirect("pregoes:relatorio_distribuicao")

            pregao_permitido = QuantitativoEscola.objects.filter(
                pregao_id=pregao_id,
                escola=escola_consulta,
                quantidade__gt=0,
            ).exists()

            if not pregao_permitido:
                messages.error(
                    request,
                    "Você não tem permissão para consultar relatórios deste pregão.",
                )
                return redirect("pregoes:relatorio_distribuicao")

        pregao = get_object_or_404(Pregao, id=pregao_id)

        itens_pregao = (
            QuantitativoPregao.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("item__nome_item")
        )

        for qp in itens_pregao:
            item = qp.item
            total_pregao = qp.quantidade

            distribuicoes_todas = (
                QuantitativoEscola.objects.filter(
                    pregao=pregao,
                    item=item,
                )
                .select_related("escola", "escola__municipio")
                .order_by("escola__nome_escola")
            )

            total_distribuido = distribuicoes_todas.aggregate(
                total=Sum("quantidade")
            ).get("total")

            if total_distribuido is None:
                total_distribuido = Decimal("0")

            # Para Consulta/Escola, exibe somente a linha da escola vinculada.
            # O total/status do item continua considerando a distribuição geral do pregão,
            # mas a lista detalhada não mostra outras escolas.
            if consulta_escola and escola_consulta:
                distribuicoes = distribuicoes_todas.filter(escola=escola_consulta)

                # Se a escola não possui este item, não mostra o item no relatório dela.
                if not distribuicoes.exists():
                    continue
            else:
                distribuicoes = distribuicoes_todas

            saldo = total_pregao - total_distribuido

            total_itens += 1

            if saldo == 0:
                status = "Fechado"
                status_classe = "badge-success"
                status_linha = "status-fechado"
                itens_fechados += 1
            elif saldo > 0:
                status = "Com saldo"
                status_classe = "badge-warning"
                status_linha = "status-saldo"
                itens_com_saldo += 1
            else:
                status = "Excedido"
                status_classe = "badge-danger"
                status_linha = "status-excedido"
                itens_excedidos += 1

            relatorio.append(
                {
                    "item": item,
                    "total_pregao": total_pregao,
                    "distribuicoes": distribuicoes,
                    "total_distribuido": total_distribuido,
                    "saldo": saldo,
                    "status": status,
                    "status_classe": status_classe,
                    "status_linha": status_linha,
                }
            )

    return render(
        request,
        "pregoes/relatorio_distribuicao.html",
        {
            "pregoes": pregoes_lista,
            "pregao": pregao,
            "relatorio": relatorio,
            "total_itens": total_itens,
            "itens_fechados": itens_fechados,
            "itens_com_saldo": itens_com_saldo,
            "itens_excedidos": itens_excedidos,
            "consulta_escola": consulta_escola,
            "escola_consulta": escola_consulta,
        },
    )

def escolas_por_pregao(request, pregao_id):
    pregao = get_object_or_404(Pregao, id=pregao_id)

    municipios_ids = pregao.municipios.values_list("id", flat=True)

    escolas = (
        Escola.objects.filter(
            ativo=True,
            municipio_id__in=municipios_ids,
        )
        .select_related("municipio")
        .order_by("nome_escola")
    )

    dados = []

    for escola in escolas:
        dados.append(
            {
                "id": escola.id,
                "nome": escola.nome_escola,
                "municipio": escola.municipio.nome if escola.municipio else "",
                "uf": escola.municipio.uf if escola.municipio and escola.municipio.uf else "",
            }
        )

    return JsonResponse({"escolas": dados})

def itens_por_pregao(request, pregao_id):
    pregao = get_object_or_404(
        Pregao,
        id=pregao_id,
        tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
    )

    # Se ainda não houver PregaoItem, tenta criar a partir dos quantitativos,
    # mas somente se já existir quantitativo por escola.
    if not PregaoItem.objects.filter(pregao=pregao).exists():
        garantir_itens_pregao_para_propostas(pregao)

    itens = (
        PregaoItem.objects.filter(pregao=pregao)
        .select_related("item")
        .order_by("ordem", "item__nome_item")
    )

    dados = []

    for item_pregao in itens:
        dados.append(
            {
                "id": item_pregao.id,
                "ordem": item_pregao.ordem,
                "nome": item_pregao.item.nome_item,
            }
        )

    return JsonResponse({"itens": dados})

def garantir_itens_pregao_para_propostas(pregao):
    """
    Para pregões ainda não iniciados, pode não existir PregaoItem.
    Esta função cria os PregaoItem a partir do QuantitativoPregao,
    mas somente se já existir pelo menos um QuantitativoEscola para o pregão.
    """

    existe_quantitativo_escola = QuantitativoEscola.objects.filter(
        pregao=pregao
    ).exists()

    if not existe_quantitativo_escola:
        return False

    quantitativos_pregao = (
        QuantitativoPregao.objects.filter(pregao=pregao)
        .select_related("item")
        .order_by("item__nome_item")
    )

    ordem = 1

    for qp in quantitativos_pregao:
        PregaoItem.objects.get_or_create(
            pregao=pregao,
            item=qp.item,
            defaults={
                "ordem": ordem,
                "quantidade_total": qp.quantidade,
            },
        )

        ordem += 1

    return True


def media_precos(request):
    pregoes_lista = (
        Pregao.objects.filter(
            status__in=[
                Pregao.STATUS_NAO_INICIADO,
                Pregao.STATUS_EM_ANDAMENTO,
            ]
        )
        .prefetch_related("municipios")
        .order_by("-ano", "-numero")
    )

    pregao_id = request.GET.get("pregao") or request.POST.get("pregao")
    pregao = None
    itens_pregao = []

    if pregao_id:
        pregao = get_object_or_404(Pregao, id=pregao_id)

        if pregao.status not in [
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]:
            messages.error(
                request,
                "Só é possível cadastrar médias de preços em certames não iniciados ou em andamento.",
            )
            return redirect("pregoes:media_precos")

        if not PregaoItem.objects.filter(pregao=pregao).exists():
            garantir_itens_pregao_para_propostas(pregao)

        itens_pregao = (
            PregaoItem.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("ordem", "item__nome_item")
        )

    if request.method == "POST":
        if not pregao:
            messages.error(request, "Selecione um certame.")
            return redirect("pregoes:media_precos")

        if pregao.status not in [
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]:
            messages.error(
                request,
                "Só é possível alterar médias de preços em certames não iniciados ou em andamento.",
            )
            return redirect("pregoes:media_precos")

        erros = []
        atualizados = 0

        # Percentual único do certame
        percentual_texto = request.POST.get("percentual_alerta_media", "").strip()

        if percentual_texto:
            percentual_texto = percentual_texto.replace(",", ".")

            try:
                percentual_alerta = Decimal(percentual_texto)
            except InvalidOperation:
                erros.append("O percentual de alerta do certame é inválido.")
                percentual_alerta = None

            if percentual_alerta is not None:
                if percentual_alerta < 0:
                    erros.append("O percentual de alerta do certame não pode ser negativo.")
                elif percentual_alerta > 100:
                    erros.append("O percentual de alerta do certame não pode ser maior que 100%.")
        else:
            percentual_alerta = Decimal("50")

        # Médias continuam sendo individuais por item
        for item_pregao in itens_pregao:
            media_texto = request.POST.get(f"media_preco_{item_pregao.id}", "").strip()
            media_preco = None

            if media_texto:
                media_texto = media_texto.replace(".", "").replace(",", ".")

                try:
                    media_preco = Decimal(media_texto)
                except InvalidOperation:
                    erros.append(
                        f"A média de preço do item '{item_pregao.item.nome_item}' é inválida."
                    )
                    continue

                if media_preco <= 0:
                    erros.append(
                        f"A média de preço do item '{item_pregao.item.nome_item}' deve ser maior que zero."
                    )
                    continue

            item_pregao.media_preco = media_preco
            item_pregao.save(update_fields=["media_preco"])
            atualizados += 1

        if erros:
            for erro in erros:
                messages.error(request, erro)

            return redirect(f"{request.path}?pregao={pregao.id}")

        pregao.percentual_alerta_media = percentual_alerta
        pregao.save(update_fields=["percentual_alerta_media", "atualizado_em"])

        messages.success(
            request,
            f"Médias de preços salvas com sucesso. Itens atualizados: {atualizados}. "
            f"Percentual de alerta do certame: {percentual_alerta:.2f}%.",
        )

        return redirect(f"{request.path}?pregao={pregao.id}")

    return render(
        request,
        "pregoes/media_precos.html",
        {
            "pregoes": pregoes_lista,
            "pregao": pregao,
            "itens_pregao": itens_pregao,
        },
    )



# ============================================================
# PROPOSTA DIGITAL - ETAPA 1
# ============================================================

PROPOSTA_DIGITAL_SIGNING_SALT = "siga-proposta-digital-cnpj-v1"
PROPOSTA_DIGITAL_CONSULTA_MAX_AGE = 12 * 60 * 60
PROPOSTA_DIGITAL_MAX_PACOTE_B64 = 2_000_000
PROPOSTA_DIGITAL_IMPORTACAO_SALT = "siga-proposta-digital-importacao-v1"
PROPOSTA_DIGITAL_IMPORTACAO_MAX_AGE = 2 * 60 * 60
PROPOSTA_DIGITAL_ALGORITMO_V1 = "AES-256-GCM/PBKDF2-SHA256/200000"
PROPOSTA_DIGITAL_ALFABETO_CODIGO = set("ABCDEFGHJKLMNPQRSTUVWXYZ23456789")


def _normalizar_codigo_abertura_proposta_digital(valor):
    texto = str(valor or "").strip().upper()

    # Permite também colar o conteúdo integral do QR Code:
    # SIGA-PROPOSTA|1|PX7K-....
    if "|" in texto:
        texto = texto.split("|")[-1]

    codigo = re.sub(r"[^A-Z0-9]", "", texto)
    if (
        len(codigo) != 20
        or any(caractere not in PROPOSTA_DIGITAL_ALFABETO_CODIGO for caractere in codigo)
    ):
        return ""

    return codigo


def _perfil_usuario_atual(user):
    if not user.is_authenticated:
        return None
    if user.is_superuser:
        return "administrador"
    perfil_acesso = getattr(user, "perfil_acesso", None)
    return getattr(perfil_acesso, "perfil", None)


def _somente_digitos(valor):
    return re.sub(r"\D", "", str(valor or ""))


def _formatar_cnpj_publico(valor):
    cnpj = _somente_digitos(valor)
    if len(cnpj) != 14:
        return str(valor or "")
    return (
        f"{cnpj[0:2]}.{cnpj[2:5]}.{cnpj[5:8]}/"
        f"{cnpj[8:12]}-{cnpj[12:14]}"
    )


def _json_no_store(payload, status=200):
    response = JsonResponse(payload, status=status)
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    return response


def _link_proposta_digital_aberto(link):
    return (
        link.ativo
        and link.pregao.tipo_certame == Pregao.TIPO_PREGAO_PRESENCIAL
        and link.pregao.status == Pregao.STATUS_NAO_INICIADO
    )


def _hash_estrutura_itens_proposta_digital(pregao):
    partes = []
    for item in (
        PregaoItem.objects.filter(pregao=pregao)
        .values_list("id", "item_id", "ordem", "quantidade_total")
        .order_by("ordem", "id")
    ):
        partes.append("|".join(str(valor) for valor in item))
    texto = ";".join(partes)
    return hashlib.sha256(texto.encode("utf-8")).hexdigest()


@login_required
@require_POST
def gerar_link_proposta_digital(request):
    """
    Cria (ou reutiliza) um único link público por Pregão Presencial.
    O mesmo link pode ser enviado a todos os fornecedores.
    """
    perfil = _perfil_usuario_atual(request.user)
    if perfil not in {"administrador", "pregoeiro"}:
        return _json_no_store(
            {"ok": False, "mensagem": "Você não tem permissão para gerar este link."},
            status=403,
        )

    pregao_id = request.POST.get("pregao_id")
    if not pregao_id:
        return _json_no_store(
            {"ok": False, "mensagem": "Selecione o pregão."},
            status=400,
        )

    pregao = get_object_or_404(
        Pregao,
        id=pregao_id,
        tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
    )

    if pregao.status != Pregao.STATUS_NAO_INICIADO:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "O link de proposta digital só pode ser gerado enquanto "
                    "o pregão estiver com status Não iniciado."
                ),
            },
            status=400,
        )

    if not PregaoItem.objects.filter(pregao=pregao).exists():
        garantir_itens_pregao_para_propostas(pregao)

    quantidade_itens = PregaoItem.objects.filter(pregao=pregao).count()
    if quantidade_itens == 0:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "Este pregão ainda não possui itens disponíveis para a proposta. "
                    "Confira os quantitativos do pregão/escolas antes de gerar o link."
                ),
            },
            status=400,
        )

    link, criado = PropostaDigitalLink.objects.get_or_create(
        pregao=pregao,
        defaults={
            "criado_por": request.user,
            "ativo": True,
        },
    )

    if not link.ativo:
        link.ativo = True
        link.save(update_fields=["ativo", "atualizado_em"])

    url_publica = request.build_absolute_uri(
        reverse("pregoes:proposta_digital_publica", args=[link.token])
    )

    return _json_no_store(
        {
            "ok": True,
            "criado": criado,
            "url": url_publica,
            "pregao": f"{pregao.numero}/{pregao.ano}",
            "quantidade_itens": quantidade_itens,
        }
    )


@never_cache
def proposta_digital_publica(request, token):
    link = get_object_or_404(
        PropostaDigitalLink.objects.select_related("pregao"),
        token=token,
        ativo=True,
    )
    pregao = link.pregao
    disponivel = _link_proposta_digital_aberto(link)

    itens_pregao = []
    if disponivel:
        itens_pregao = list(
            PregaoItem.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("ordem", "item__nome_item")
        )

    response = render(
        request,
        "pregoes/proposta_digital_publica.html",
        {
            "link_proposta": link,
            "pregao": pregao,
            "itens_pregao": itens_pregao,
            "disponivel": disponivel and bool(itens_pregao),
            "estrutura_hash": (
                _hash_estrutura_itens_proposta_digital(pregao)
                if disponivel and itens_pregao
                else ""
            ),
        },
    )
    response["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response["Pragma"] = "no-cache"
    return response


@never_cache
@require_GET
def consultar_cnpj_proposta_digital(request, token):
    """
    Consulta apenas CNPJ + Razão Social na BrasilAPI para conferência do fornecedor.
    Não consulta cadastro de fornecedores nem vínculo com o pregão.
    """
    link = get_object_or_404(
        PropostaDigitalLink.objects.select_related("pregao"),
        token=token,
        ativo=True,
    )

    if not _link_proposta_digital_aberto(link):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Este link não está mais aberto para recebimento de propostas.",
            },
            status=410,
        )

    cnpj = _somente_digitos(request.GET.get("cnpj"))
    if len(cnpj) != 14:
        return _json_no_store(
            {"ok": False, "mensagem": "Informe um CNPJ com 14 dígitos."},
            status=400,
        )

    try:
        validar_cnpj(cnpj)
    except ValidationError:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "CNPJ inválido. Confira os números informados.",
            },
            status=400,
        )

    # Antes de consultar qualquer serviço externo, verifica se este CNPJ
    # já possui uma proposta digital registrada para o mesmo pregão.
    # O CNPJ continua sem ser armazenado em texto puro: a comparação é feita
    # somente pelo SHA-256, exatamente como na finalização da proposta.
    cnpj_hash_consulta = hashlib.sha256(cnpj.encode("utf-8")).hexdigest()

    if PropostaDigital.objects.filter(
        pregao=link.pregao,
        cnpj_hash=cnpj_hash_consulta,
        status__in=[
            PropostaDigital.STATUS_FINALIZADA,
            PropostaDigital.STATUS_IMPORTADA,
        ],
    ).exists():
        return _json_no_store(
            {
                "ok": False,
                "proposta_existente": True,
                "permitir_razao_manual": False,
                "mensagem": (
                    "Já existe uma proposta finalizada para este CNPJ neste pregão. "
                    "Não é necessário preencher uma nova proposta."
                ),
            },
            status=409,
        )

    razao_social_manual = " ".join(
        str(request.GET.get("razao_social_manual") or "").strip().split()
    )

    # Consulta com redundância:
    # 1) BrasilAPI
    # 2) Minha Receita
    # 3) CNPJ.ws (API pública), como fallback adicional.
    #
    # Não há qualquer consulta ao cadastro de fornecedores do SIGA aqui.
    provedores = [
        (
            "BrasilAPI",
            f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}",
        ),
        (
            "Minha Receita",
            f"https://minhareceita.org/{cnpj}",
        ),
        (
            "CNPJ.ws",
            f"https://publica.cnpj.ws/cnpj/{cnpj}",
        ),
    ]

    dados = None
    erros_provedores = []

    for nome_provedor, url in provedores:
        requisicao = Request(
            url,
            headers={
                "Accept": "application/json",
                "User-Agent": (
                    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                    "AppleWebKit/537.36 (KHTML, like Gecko) "
                    "Chrome/154.0 Safari/537.36 SIGA-Escolar/1.0"
                ),
            },
        )

        try:
            with urlopen(requisicao, timeout=12) as resposta:
                candidato = json.loads(resposta.read().decode("utf-8"))

            # Os provedores utilizados devolvem "razao_social" no nível
            # principal. Se houver resposta 200 sem esse campo, o SIGA tenta
            # automaticamente o próximo serviço.
            if isinstance(candidato, dict) and str(
                candidato.get("razao_social") or ""
            ).strip():
                dados = candidato
                break

            erros_provedores.append(f"{nome_provedor}: resposta incompleta")

        except HTTPError as erro:
            erros_provedores.append(f"{nome_provedor}: HTTP {erro.code}")
            continue

        except (URLError, TimeoutError, json.JSONDecodeError) as erro:
            erros_provedores.append(
                f"{nome_provedor}: falha de conexão/resposta"
            )
            continue

    if dados is None:
        # Se todos os serviços responderam 404, não tratamos isso como
        # indisponibilidade técnica: o CNPJ simplesmente não foi localizado.
        if erros_provedores and all("HTTP 404" in erro for erro in erros_provedores):
            return _json_no_store(
                {
                    "ok": False,
                    "mensagem": "CNPJ não encontrado nas bases públicas de consulta.",
                    "permitir_razao_manual": False,
                },
                status=404,
            )

        # A digitação manual só é aceita neste ponto: isto é, depois de as
        # três fontes externas terem sido tentadas sem retornar Razão Social.
        if razao_social_manual:
            if len(razao_social_manual) < 3:
                return _json_no_store(
                    {
                        "ok": False,
                        "mensagem": "Informe uma Razão Social válida.",
                        "permitir_razao_manual": True,
                        "cnpj": _formatar_cnpj_publico(cnpj),
                    },
                    status=400,
                )

            if len(razao_social_manual) > 255:
                return _json_no_store(
                    {
                        "ok": False,
                        "mensagem": "A Razão Social deve possuir no máximo 255 caracteres.",
                        "permitir_razao_manual": True,
                        "cnpj": _formatar_cnpj_publico(cnpj),
                    },
                    status=400,
                )

            consulta_token = signing.dumps(
                {
                    "pregao_id": link.pregao_id,
                    "cnpj": cnpj,
                    "razao_social": razao_social_manual,
                    "origem_razao_social": "manual",
                },
                salt=PROPOSTA_DIGITAL_SIGNING_SALT,
                compress=True,
            )

            return _json_no_store(
                {
                    "ok": True,
                    "dados": {
                        "cnpj": _formatar_cnpj_publico(cnpj),
                        "razao_social": razao_social_manual,
                    },
                    "consulta_token": consulta_token,
                    "origem_razao_social": "manual",
                    "aviso": (
                        "Razão Social informada manualmente porque os serviços "
                        "públicos de consulta estavam indisponíveis."
                    ),
                }
            )

        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "Os serviços públicos de consulta de CNPJ estão indisponíveis "
                    "no momento. Informe a Razão Social manualmente para continuar."
                ),
                "permitir_razao_manual": True,
                "cnpj": _formatar_cnpj_publico(cnpj),
            },
            status=502,
        )

    razao_social = str(dados.get("razao_social") or "").strip()
    if not razao_social:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "A consulta não retornou a Razão Social deste CNPJ.",
            },
            status=502,
        )

    # Token assinado e temporário: comprova no POST final que este CNPJ foi
    # efetivamente consultado em fonte pública, sem exigir cadastro prévio no SIGA.
    consulta_token = signing.dumps(
        {
            "pregao_id": link.pregao_id,
            "cnpj": cnpj,
            "razao_social": razao_social,
            "origem_razao_social": "consulta_publica",
        },
        salt=PROPOSTA_DIGITAL_SIGNING_SALT,
        compress=True,
    )

    return _json_no_store(
        {
            "ok": True,
            "dados": {
                "cnpj": _formatar_cnpj_publico(cnpj),
                "razao_social": razao_social,
            },
            "consulta_token": consulta_token,
            "origem_razao_social": "consulta_publica",
        }
    )


@never_cache
@require_POST
def finalizar_proposta_digital(request, token):
    link = get_object_or_404(
        PropostaDigitalLink.objects.select_related("pregao"),
        token=token,
        ativo=True,
    )
    pregao = link.pregao

    if not _link_proposta_digital_aberto(link):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Este link não está mais aberto para recebimento de propostas.",
            },
            status=410,
        )

    try:
        dados = json.loads(request.body.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError):
        return _json_no_store(
            {"ok": False, "mensagem": "Dados da proposta inválidos."},
            status=400,
        )

    consulta_token = str(dados.get("consulta_token") or "")
    try:
        consulta = signing.loads(
            consulta_token,
            salt=PROPOSTA_DIGITAL_SIGNING_SALT,
            max_age=PROPOSTA_DIGITAL_CONSULTA_MAX_AGE,
        )
    except signing.SignatureExpired:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "A consulta do CNPJ expirou. Consulte o CNPJ novamente "
                    "antes de finalizar a proposta."
                ),
            },
            status=400,
        )
    except signing.BadSignature:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "A validação do CNPJ não pôde ser confirmada.",
            },
            status=400,
        )

    if str(consulta.get("pregao_id")) != str(pregao.id):
        return _json_no_store(
            {"ok": False, "mensagem": "A consulta do CNPJ pertence a outro pregão."},
            status=400,
        )

    cnpj = _somente_digitos(consulta.get("cnpj"))
    cnpj_hash_esperado = hashlib.sha256(cnpj.encode("utf-8")).hexdigest()
    cnpj_hash = str(dados.get("cnpj_hash") or "").lower()

    if not hmac.compare_digest(cnpj_hash, cnpj_hash_esperado):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "O CNPJ da proposta não corresponde ao CNPJ consultado.",
            },
            status=400,
        )

    codigo_hash = str(dados.get("codigo_hash") or "").lower()
    conteudo_hash = str(dados.get("conteudo_hash") or "").lower()
    padrao_hash = re.compile(r"^[0-9a-f]{64}$")

    if not padrao_hash.fullmatch(codigo_hash):
        return _json_no_store(
            {"ok": False, "mensagem": "Código de abertura inválido."},
            status=400,
        )
    if not padrao_hash.fullmatch(conteudo_hash):
        return _json_no_store(
            {"ok": False, "mensagem": "Hash de integridade inválido."},
            status=400,
        )

    pacote_b64 = str(dados.get("pacote_criptografado") or "")
    iv_b64 = str(dados.get("iv_b64") or "")
    salt_b64 = str(dados.get("salt_b64") or "")

    if not pacote_b64 or len(pacote_b64) > PROPOSTA_DIGITAL_MAX_PACOTE_B64:
        return _json_no_store(
            {"ok": False, "mensagem": "Tamanho do pacote da proposta inválido."},
            status=400,
        )

    try:
        pacote_bytes = base64.b64decode(pacote_b64, validate=True)
        iv_bytes = base64.b64decode(iv_b64, validate=True)
        salt_bytes = base64.b64decode(salt_b64, validate=True)
    except (ValueError, TypeError, binascii.Error):
        return _json_no_store(
            {"ok": False, "mensagem": "Pacote criptografado inválido."},
            status=400,
        )

    if len(iv_bytes) != 12 or len(salt_bytes) != 16 or len(pacote_bytes) <= 16:
        return _json_no_store(
            {"ok": False, "mensagem": "Parâmetros criptográficos inválidos."},
            status=400,
        )

    hash_real = hashlib.sha256(pacote_bytes).hexdigest()
    if not hmac.compare_digest(hash_real, conteudo_hash):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Falha na verificação de integridade do pacote criptografado.",
            },
            status=400,
        )

    try:
        quantidade_itens = int(dados.get("quantidade_itens"))
    except (TypeError, ValueError):
        quantidade_itens = -1

    quantidade_esperada = PregaoItem.objects.filter(pregao=pregao).count()
    estrutura_hash = str(dados.get("estrutura_hash") or "").lower()
    estrutura_atual = _hash_estrutura_itens_proposta_digital(pregao)
    if (
        quantidade_itens != quantidade_esperada
        or quantidade_esperada <= 0
        or not padrao_hash.fullmatch(estrutura_hash)
        or not hmac.compare_digest(estrutura_hash, estrutura_atual)
    ):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "A relação de itens do pregão foi alterada. Atualize a página "
                    "e confira a proposta antes de finalizar."
                ),
            },
            status=409,
        )

    # Finalização idempotente:
    # se o navegador não recebeu a resposta anterior (queda de conexão,
    # proxy retornando HTML, etc.), ele pode reenviar exatamente a mesma
    # tentativa. Como o servidor só possui o hash do código, o mesmo
    # codigo_hash identifica com segurança a mesma abertura.
    proposta_mesmo_codigo = PropostaDigital.objects.filter(
        codigo_hash=codigo_hash,
    ).first()

    if proposta_mesmo_codigo:
        if (
            proposta_mesmo_codigo.pregao_id == pregao.id
            and hmac.compare_digest(
                proposta_mesmo_codigo.cnpj_hash,
                cnpj_hash,
            )
        ):
            return _json_no_store(
                {
                    "ok": True,
                    "protocolo": str(proposta_mesmo_codigo.protocolo),
                    "mensagem": "Proposta registrada com sucesso.",
                    "ja_registrada": True,
                },
                status=200,
            )

        return _json_no_store(
            {
                "ok": False,
                "mensagem": "O código de abertura informado não pode ser utilizado.",
            },
            status=409,
        )

    proposta_mesmo_cnpj = PropostaDigital.objects.filter(
        pregao=pregao,
        cnpj_hash=cnpj_hash,
        status__in=[
            PropostaDigital.STATUS_FINALIZADA,
            PropostaDigital.STATUS_IMPORTADA,
        ],
    ).first()

    if proposta_mesmo_cnpj:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "Já existe uma proposta finalizada para este CNPJ neste pregão. "
                    "Após a finalização, uma nova proposta não pode ser enviada pelo mesmo CNPJ."
                ),
            },
            status=409,
        )

    try:
        proposta = PropostaDigital.objects.create(
            pregao=pregao,
            link=link,
            codigo_hash=codigo_hash,
            cnpj_hash=cnpj_hash,
            pacote_criptografado=pacote_b64,
            iv_b64=iv_b64,
            salt_b64=salt_b64,
            conteudo_hash=conteudo_hash,
            estrutura_hash=estrutura_hash,
            quantidade_itens=quantidade_itens,
        )
    except IntegrityError:
        # Proteção contra concorrência: outra requisição pode ter concluído
        # entre as verificações acima e o INSERT.
        proposta_mesmo_codigo = PropostaDigital.objects.filter(
            codigo_hash=codigo_hash,
        ).first()

        if (
            proposta_mesmo_codigo
            and proposta_mesmo_codigo.pregao_id == pregao.id
            and hmac.compare_digest(
                proposta_mesmo_codigo.cnpj_hash,
                cnpj_hash,
            )
        ):
            return _json_no_store(
                {
                    "ok": True,
                    "protocolo": str(proposta_mesmo_codigo.protocolo),
                    "mensagem": "Proposta registrada com sucesso.",
                    "ja_registrada": True,
                },
                status=200,
            )

        if PropostaDigital.objects.filter(
            pregao=pregao,
            cnpj_hash=cnpj_hash,
            status__in=[
                PropostaDigital.STATUS_FINALIZADA,
                PropostaDigital.STATUS_IMPORTADA,
            ],
        ).exists():
            mensagem = (
                "Já existe uma proposta finalizada para este CNPJ neste pregão. "
                "Após a finalização, uma nova proposta não pode ser enviada pelo mesmo CNPJ."
            )
        else:
            mensagem = "Não foi possível registrar a proposta. Tente novamente sem fechar esta página."

        return _json_no_store(
            {"ok": False, "mensagem": mensagem},
            status=409,
        )

    return _json_no_store(
        {
            "ok": True,
            "protocolo": str(proposta.protocolo),
            "mensagem": "Proposta registrada com sucesso.",
            "ja_registrada": False,
        },
        status=201,
    )


@login_required
@require_POST
def localizar_proposta_digital(request):
    """
    Localiza uma Proposta Digital pelo código entregue pelo fornecedor.

    O servidor:
    - usa somente o hash do código para localizar a proposta;
    - confere Pregão + CNPJ do fornecedor selecionado;
    - não descriptografa marca ou preços;
    - devolve o pacote criptografado para abertura no navegador.

    A proposta somente é marcada como importada quando o usuário confirmar
    o salvamento das Propostas Iniciais.
    """
    perfil = _perfil_usuario_atual(request.user)
    if perfil not in {"administrador", "pregoeiro"}:
        return _json_no_store(
            {"ok": False, "mensagem": "Você não tem permissão para importar propostas digitais."},
            status=403,
        )

    pregao_id = request.POST.get("pregao_id")
    fornecedor_id = request.POST.get("fornecedor_id")
    codigo = _normalizar_codigo_abertura_proposta_digital(
        request.POST.get("codigo")
    )

    if not pregao_id or not fornecedor_id:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Selecione o pregão e o fornecedor antes de informar o código.",
            },
            status=400,
        )

    if not codigo:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Código inválido. Confira o código fornecido pelo fornecedor.",
            },
            status=400,
        )

    pregao = get_object_or_404(
        Pregao,
        id=pregao_id,
        tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
    )

    if pregao.status not in [
        Pregao.STATUS_NAO_INICIADO,
        Pregao.STATUS_EM_ANDAMENTO,
    ]:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Este pregão não está disponível para importação de propostas.",
            },
            status=409,
        )

    vinculo_fornecedor = (
        PregaoFornecedor.objects.filter(
            pregao=pregao,
            fornecedor_id=fornecedor_id,
            ativo_no_pregao=True,
        )
        .select_related("fornecedor")
        .first()
    )

    if not vinculo_fornecedor:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "O fornecedor selecionado não está ativo neste pregão.",
            },
            status=400,
        )

    fornecedor = vinculo_fornecedor.fornecedor
    cnpj_fornecedor = _somente_digitos(fornecedor.cnpj)

    if len(cnpj_fornecedor) != 14:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "O fornecedor selecionado não possui um CNPJ válido cadastrado. "
                    "Confira o cadastro antes de importar a proposta."
                ),
            },
            status=400,
        )

    codigo_hash = hashlib.sha256(codigo.encode("utf-8")).hexdigest()

    proposta = (
        PropostaDigital.objects.select_related("pregao")
        .filter(codigo_hash=codigo_hash)
        .first()
    )

    # Não revelamos a qual pregão ou fornecedor pertence um código que não confere.
    mensagem_nao_confere = (
        "O código informado não corresponde ao fornecedor selecionado "
        "ou ao pregão atual. Nenhuma informação foi importada."
    )

    if not proposta or proposta.pregao_id != pregao.id:
        return _json_no_store(
            {"ok": False, "mensagem": mensagem_nao_confere},
            status=404,
        )

    cnpj_hash_fornecedor = hashlib.sha256(
        cnpj_fornecedor.encode("utf-8")
    ).hexdigest()

    if not hmac.compare_digest(
        proposta.cnpj_hash,
        cnpj_hash_fornecedor,
    ):
        return _json_no_store(
            {"ok": False, "mensagem": mensagem_nao_confere},
            status=409,
        )

    if proposta.status == PropostaDigital.STATUS_IMPORTADA:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Esta Proposta Digital já foi importada anteriormente.",
            },
            status=409,
        )

    if proposta.status != PropostaDigital.STATUS_FINALIZADA:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Esta Proposta Digital não está disponível para importação.",
            },
            status=409,
        )

    if proposta.algoritmo != PROPOSTA_DIGITAL_ALGORITMO_V1:
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "Esta proposta utiliza uma versão de criptografia não suportada "
                    "por esta versão do SIGA."
                ),
            },
            status=409,
        )

    estrutura_atual = _hash_estrutura_itens_proposta_digital(pregao)
    quantidade_atual = PregaoItem.objects.filter(pregao=pregao).count()

    if (
        proposta.quantidade_itens != quantidade_atual
        or not hmac.compare_digest(
            proposta.estrutura_hash,
            estrutura_atual,
        )
    ):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "A relação de itens deste pregão foi alterada após o envio "
                    "da proposta. A importação foi bloqueada para conferência."
                ),
            },
            status=409,
        )

    try:
        pacote_bytes = base64.b64decode(
            proposta.pacote_criptografado,
            validate=True,
        )
        iv_bytes = base64.b64decode(proposta.iv_b64, validate=True)
        salt_bytes = base64.b64decode(proposta.salt_b64, validate=True)
    except (ValueError, TypeError, binascii.Error):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "O pacote criptografado desta proposta está inválido.",
            },
            status=409,
        )

    if (
        len(iv_bytes) != 12
        or len(salt_bytes) != 16
        or len(pacote_bytes) <= 16
        or len(proposta.pacote_criptografado) > PROPOSTA_DIGITAL_MAX_PACOTE_B64
    ):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": "Os parâmetros criptográficos desta proposta estão inválidos.",
            },
            status=409,
        )

    conteudo_hash_real = hashlib.sha256(pacote_bytes).hexdigest()
    if not hmac.compare_digest(
        conteudo_hash_real,
        proposta.conteudo_hash,
    ):
        return _json_no_store(
            {
                "ok": False,
                "mensagem": (
                    "A verificação de integridade da Proposta Digital falhou. "
                    "Nenhuma informação foi importada."
                ),
            },
            status=409,
        )

    token_importacao = signing.dumps(
        {
            "proposta_digital_id": proposta.id,
            "pregao_id": pregao.id,
            "fornecedor_id": fornecedor.id,
            "codigo_hash": proposta.codigo_hash,
            "estrutura_hash": proposta.estrutura_hash,
        },
        salt=PROPOSTA_DIGITAL_IMPORTACAO_SALT,
        compress=True,
    )

    return _json_no_store(
        {
            "ok": True,
            "protocolo": str(proposta.protocolo),
            "pacote_criptografado": proposta.pacote_criptografado,
            "iv_b64": proposta.iv_b64,
            "salt_b64": proposta.salt_b64,
            "conteudo_hash": proposta.conteudo_hash,
            "estrutura_hash": proposta.estrutura_hash,
            "quantidade_itens": proposta.quantidade_itens,
            "algoritmo": proposta.algoritmo,
            "token_importacao": token_importacao,
        }
    )



@login_required
def controle_propostas_digitais(request):
    """Acompanha Propostas Digitais sem quebrar o sigilo antes da importação."""
    perfil = _perfil_usuario_atual(request.user)
    if perfil not in {"administrador", "pregoeiro"}:
        messages.error(
            request,
            "Você não tem permissão para acessar o Controle de Propostas Digitais.",
        )
        return redirect("pregoes:propostas_iniciais")

    pregoes_lista = (
        Pregao.objects.filter(tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL)
        .prefetch_related("municipios")
        .order_by("-ano", "-numero")
    )

    pregao_id = request.GET.get("pregao")
    pregao = None
    propostas = PropostaDigital.objects.none()
    resumo = {"total": 0, "finalizadas": 0, "importadas": 0, "invalidadas": 0}

    if pregao_id:
        pregao = get_object_or_404(
            Pregao,
            id=pregao_id,
            tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
        )
        propostas = (
            PropostaDigital.objects.filter(pregao=pregao)
            .select_related(
                "fornecedor_importado",
                "importado_por",
                "invalidada_por",
            )
            .order_by("-finalizada_em", "-id")
        )
        resumo = {
            "total": propostas.count(),
            "finalizadas": propostas.filter(
                status=PropostaDigital.STATUS_FINALIZADA
            ).count(),
            "importadas": propostas.filter(
                status=PropostaDigital.STATUS_IMPORTADA
            ).count(),
            "invalidadas": propostas.filter(
                status=PropostaDigital.STATUS_INVALIDADA
            ).count(),
        }

    return render(
        request,
        "pregoes/controle_propostas_digitais.html",
        {
            "pregoes": pregoes_lista,
            "pregao": pregao,
            "propostas": propostas,
            "resumo": resumo,
        },
    )


@login_required
@require_POST
def invalidar_proposta_digital(request, proposta_id):
    """Invalida uma proposta ainda lacrada, preservando o histórico."""
    perfil = _perfil_usuario_atual(request.user)
    if perfil not in {"administrador", "pregoeiro"}:
        messages.error(
            request,
            "Você não tem permissão para invalidar Propostas Digitais.",
        )
        return redirect("pregoes:controle_propostas_digitais")

    motivo = " ".join(str(request.POST.get("motivo") or "").strip().split())
    pregao_retorno = request.POST.get("pregao_id")

    if len(motivo) < 10:
        messages.error(
            request,
            "Informe um motivo para a invalidação com pelo menos 10 caracteres.",
        )
        if pregao_retorno:
            return redirect(
                f"{reverse('pregoes:controle_propostas_digitais')}?pregao={pregao_retorno}"
            )
        return redirect("pregoes:controle_propostas_digitais")

    if len(motivo) > 1000:
        messages.error(
            request,
            "O motivo da invalidação deve possuir no máximo 1000 caracteres.",
        )
        if pregao_retorno:
            return redirect(
                f"{reverse('pregoes:controle_propostas_digitais')}?pregao={pregao_retorno}"
            )
        return redirect("pregoes:controle_propostas_digitais")

    try:
        with transaction.atomic():
            proposta = (
                PropostaDigital.objects.select_for_update()
                .select_related("pregao")
                .filter(id=proposta_id)
                .first()
            )

            if not proposta:
                raise ValidationError("A Proposta Digital não foi encontrada.")

            pregao = proposta.pregao
            pregao_retorno = pregao.id

            if pregao.status != Pregao.STATUS_NAO_INICIADO:
                raise ValidationError(
                    "A proposta só pode ser invalidada enquanto o pregão estiver "
                    "com situação Não iniciado."
                )

            if proposta.status == PropostaDigital.STATUS_IMPORTADA:
                raise ValidationError(
                    "Esta Proposta Digital já foi importada. Esta ação não desfaz "
                    "propostas iniciais já gravadas."
                )

            if proposta.status == PropostaDigital.STATUS_INVALIDADA:
                raise ValidationError("Esta Proposta Digital já está invalidada.")

            if proposta.status != PropostaDigital.STATUS_FINALIZADA:
                raise ValidationError(
                    "Esta Proposta Digital não está disponível para invalidação."
                )

            proposta.status = PropostaDigital.STATUS_INVALIDADA
            proposta.motivo_invalidacao = motivo
            proposta.invalidada_por = request.user
            proposta.invalidada_em = timezone.now()
            proposta.save(
                update_fields=[
                    "status",
                    "motivo_invalidacao",
                    "invalidada_por",
                    "invalidada_em",
                ]
            )

    except ValidationError as erro:
        messages.error(request, str(erro))
    else:
        messages.success(
            request,
            "Proposta Digital invalidada com sucesso. "
            "O fornecedor poderá enviar uma nova proposta para este pregão.",
        )

    if pregao_retorno:
        return redirect(
            f"{reverse('pregoes:controle_propostas_digitais')}?pregao={pregao_retorno}"
        )
    return redirect("pregoes:controle_propostas_digitais")

def propostas_iniciais(request):
    pregoes_lista = (
        Pregao.objects.filter(
            tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
            status__in=[
                Pregao.STATUS_NAO_INICIADO,
                Pregao.STATUS_EM_ANDAMENTO,
            ],
        )
        .prefetch_related("municipios")
        .order_by("-ano", "-numero")
    )

    modo = request.GET.get("modo") or request.POST.get("modo") or "item"
    if modo not in ["item", "fornecedor"]:
        modo = "item"

    pregao_id = request.GET.get("pregao") or request.POST.get("pregao")
    item_id = request.GET.get("item") or request.POST.get("item")
    fornecedor_id = request.GET.get("fornecedor") or request.POST.get("fornecedor")

    pregao = None
    item_pregao = None
    fornecedor_selecionado = None
    itens_pregao = []
    fornecedores = []
    propostas_salvas = {}
    pode_controlar_propostas_digitais = (
        _perfil_usuario_atual(request.user) in {"administrador", "pregoeiro"}
    )
    proposta_digital_import_token = (
        request.POST.get("proposta_digital_import_token", "").strip()
        if request.method == "POST"
        else ""
    )

    if pregao_id:
        pregao = get_object_or_404(
            Pregao,
            id=pregao_id,
            tipo_certame=Pregao.TIPO_PREGAO_PRESENCIAL,
        )

        if pregao.status not in [
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]:
            messages.error(
                request,
                "Só é possível cadastrar propostas iniciais em pregões não iniciados ou em andamento.",
            )
            return redirect("pregoes:propostas_iniciais")

        if not PregaoItem.objects.filter(pregao=pregao).exists():
            garantir_itens_pregao_para_propostas(pregao)

        itens_pregao = list(
            PregaoItem.objects.filter(pregao=pregao)
            .select_related("item")
            .order_by("ordem", "item__nome_item")
        )

        fornecedores = [
            vinculo.fornecedor
            for vinculo in (
                PregaoFornecedor.objects.filter(
                    pregao=pregao,
                    ativo_no_pregao=True,
                )
                .select_related("fornecedor")
                .order_by("ordem_inicial", "fornecedor__razao_social")
            )
        ]

    # ============================================================
    # MODO 1: CADASTRO POR ITEM
    # ============================================================
    if modo == "item" and pregao and item_id:
        item_pregao = get_object_or_404(
            PregaoItem.objects.select_related("item"),
            id=item_id,
            pregao=pregao,
        )

        propostas_salvas = {
            proposta.fornecedor_id: proposta
            for proposta in PropostaInicialItem.objects.filter(
                pregao=pregao,
                pregao_item=item_pregao,
            )
        }

    # ============================================================
    # MODO 2: CADASTRO POR FORNECEDOR
    # ============================================================
    if modo == "fornecedor" and pregao and fornecedor_id:
        fornecedor_selecionado = next(
            (
                fornecedor
                for fornecedor in fornecedores
                if str(fornecedor.id) == str(fornecedor_id)
            ),
            None,
        )

        if not fornecedor_selecionado:
            messages.error(
                request,
                "O fornecedor selecionado não está ativo neste pregão.",
            )
            return redirect(
                f"{request.path}?modo=fornecedor&pregao={pregao.id}"
            )

        propostas_salvas = {
            proposta.pregao_item_id: proposta
            for proposta in PropostaInicialItem.objects.filter(
                pregao=pregao,
                fornecedor=fornecedor_selecionado,
            )
        }

    if request.method == "POST":
        if not pregao:
            messages.error(request, "Selecione o pregão.")
            return redirect("pregoes:propostas_iniciais")

        if pregao.status not in [
            Pregao.STATUS_NAO_INICIADO,
            Pregao.STATUS_EM_ANDAMENTO,
        ]:
            messages.error(
                request,
                "Só é possível alterar propostas iniciais em pregões não iniciados ou em andamento.",
            )
            return redirect("pregoes:propostas_iniciais")

        # --------------------------------------------------------
        # SALVAR POR ITEM
        # --------------------------------------------------------
        if modo == "item":
            if not item_pregao:
                messages.error(request, "Selecione o pregão e o item.")
                return redirect(
                    f"{request.path}?modo=item&pregao={pregao.id}"
                )

            erros = []
            salvos = 0

            for fornecedor in fornecedores:
                participa = request.POST.get(
                    f"participa_{fornecedor.id}"
                ) == "on"
                marca = request.POST.get(
                    f"marca_{fornecedor.id}",
                    "",
                ).strip()
                preco_texto = request.POST.get(
                    f"preco_inicial_{fornecedor.id}",
                    "",
                ).strip()

                if not participa:
                    PropostaInicialItem.objects.update_or_create(
                        pregao=pregao,
                        pregao_item=item_pregao,
                        fornecedor=fornecedor,
                        defaults={
                            "participa": False,
                            "marca": "",
                            "preco_inicial": None,
                        },
                    )
                    salvos += 1
                    continue

                if not marca:
                    erros.append(
                        f"Informe a marca para o fornecedor "
                        f"{fornecedor.razao_social}, ou desmarque a opção Participa."
                    )
                    continue

                if not preco_texto:
                    erros.append(
                        f"Informe o preço inicial para o fornecedor "
                        f"{fornecedor.razao_social}, ou desmarque a opção Participa."
                    )
                    continue

                preco_texto = preco_texto.replace(",", ".")

                try:
                    preco_inicial = Decimal(preco_texto)
                except InvalidOperation:
                    erros.append(
                        f"Preço inicial inválido para o fornecedor "
                        f"{fornecedor.razao_social}."
                    )
                    continue

                if preco_inicial <= 0:
                    erros.append(
                        f"O preço inicial do fornecedor "
                        f"{fornecedor.razao_social} deve ser maior que zero."
                    )
                    continue

                PropostaInicialItem.objects.update_or_create(
                    pregao=pregao,
                    pregao_item=item_pregao,
                    fornecedor=fornecedor,
                    defaults={
                        "participa": True,
                        "marca": marca,
                        "preco_inicial": preco_inicial,
                    },
                )
                salvos += 1

            if erros:
                for erro in erros:
                    messages.error(request, erro)

                propostas_salvas = {
                    proposta.fornecedor_id: proposta
                    for proposta in PropostaInicialItem.objects.filter(
                        pregao=pregao,
                        pregao_item=item_pregao,
                    )
                }

                return render(
                    request,
                    "pregoes/propostas_iniciais.html",
                    {
                        "pregoes": pregoes_lista,
                        "modo": modo,
                        "pregao": pregao,
                        "itens_pregao": itens_pregao,
                        "item_pregao": item_pregao,
                        "fornecedores": fornecedores,
                        "fornecedor_selecionado": None,
                        "propostas_salvas": propostas_salvas,
                        "proposta_digital_import_token": proposta_digital_import_token,
            "pode_controlar_propostas_digitais": pode_controlar_propostas_digitais,
                    },
                )

            messages.success(
                request,
                f"Propostas iniciais salvas com sucesso. Registros salvos: {salvos}.",
            )

            acao = request.POST.get("acao")

            if acao == "salvar_avancar":
                proximo_item = (
                    PregaoItem.objects.filter(
                        pregao=pregao,
                        ordem__gt=item_pregao.ordem,
                    )
                    .select_related("item")
                    .order_by("ordem", "item__nome_item")
                    .first()
                )

                if proximo_item:
                    return redirect(
                        f"{request.path}?modo=item&pregao={pregao.id}"
                        f"&item={proximo_item.id}"
                    )

                messages.success(
                    request,
                    "Este era o último item do pregão. Não há próximo item para avançar.",
                )

            return redirect(
                f"{request.path}?modo=item&pregao={pregao.id}"
                f"&item={item_pregao.id}"
            )

        # --------------------------------------------------------
        # SALVAR POR FORNECEDOR
        # --------------------------------------------------------
        if modo == "fornecedor":
            if not fornecedor_selecionado:
                messages.error(
                    request,
                    "Selecione o pregão e o fornecedor.",
                )
                return redirect(
                    f"{request.path}?modo=fornecedor&pregao={pregao.id}"
                )

            erros = []
            dados_validados = []
            proposta_digital_para_importar_id = None

            # Se a tela foi preenchida por uma Proposta Digital, o token assinado
            # confirma que o código foi validado para este Pregão + Fornecedor.
            if proposta_digital_import_token:
                try:
                    dados_importacao = signing.loads(
                        proposta_digital_import_token,
                        salt=PROPOSTA_DIGITAL_IMPORTACAO_SALT,
                        max_age=PROPOSTA_DIGITAL_IMPORTACAO_MAX_AGE,
                    )
                except signing.SignatureExpired:
                    erros.append(
                        "A autorização da Proposta Digital expirou. "
                        "Informe o código novamente antes de salvar."
                    )
                except signing.BadSignature:
                    erros.append(
                        "Não foi possível validar a Proposta Digital importada. "
                        "Informe o código novamente."
                    )
                else:
                    if (
                        str(dados_importacao.get("pregao_id")) != str(pregao.id)
                        or str(dados_importacao.get("fornecedor_id"))
                        != str(fornecedor_selecionado.id)
                    ):
                        erros.append(
                            "A Proposta Digital importada não corresponde ao "
                            "pregão/fornecedor selecionado."
                        )
                    else:
                        proposta_digital_para_importar_id = (
                            dados_importacao.get("proposta_digital_id")
                        )
                        proposta_digital_validacao = (
                            PropostaDigital.objects.filter(
                                id=proposta_digital_para_importar_id,
                                pregao=pregao,
                                status=PropostaDigital.STATUS_FINALIZADA,
                            )
                            .first()
                        )

                        if not proposta_digital_validacao:
                            erros.append(
                                "Esta Proposta Digital não está mais disponível "
                                "para importação."
                            )
                        else:
                            cnpj_fornecedor = _somente_digitos(
                                fornecedor_selecionado.cnpj
                            )
                            cnpj_hash_fornecedor = hashlib.sha256(
                                cnpj_fornecedor.encode("utf-8")
                            ).hexdigest()

                            if (
                                len(cnpj_fornecedor) != 14
                                or not hmac.compare_digest(
                                    proposta_digital_validacao.cnpj_hash,
                                    cnpj_hash_fornecedor,
                                )
                                or not hmac.compare_digest(
                                    proposta_digital_validacao.codigo_hash,
                                    str(dados_importacao.get("codigo_hash") or ""),
                                )
                                or not hmac.compare_digest(
                                    proposta_digital_validacao.estrutura_hash,
                                    _hash_estrutura_itens_proposta_digital(pregao),
                                )
                            ):
                                erros.append(
                                    "A validação da Proposta Digital não confere. "
                                    "Informe o código novamente."
                                )

            # Primeiro valida tudo. Só grava depois que todas as linhas
            # selecionadas estiverem corretas.
            for pi in itens_pregao:
                participa = request.POST.get(
                    f"participa_item_{pi.id}"
                ) == "on"
                marca = request.POST.get(
                    f"marca_item_{pi.id}",
                    "",
                ).strip()
                preco_texto = request.POST.get(
                    f"preco_item_{pi.id}",
                    "",
                ).strip()

                if not participa:
                    dados_validados.append(
                        {
                            "pregao_item": pi,
                            "participa": False,
                            "marca": "",
                            "preco_inicial": None,
                        }
                    )
                    continue

                if not marca:
                    erros.append(
                        f"Informe a marca do item {pi.ordem} - "
                        f"{pi.item.nome_item}, ou desmarque a participação."
                    )
                    continue

                if not preco_texto:
                    erros.append(
                        f"Informe o preço inicial do item {pi.ordem} - "
                        f"{pi.item.nome_item}, ou desmarque a participação."
                    )
                    continue

                preco_normalizado = preco_texto.replace(",", ".")

                try:
                    preco_inicial = Decimal(preco_normalizado)
                except InvalidOperation:
                    erros.append(
                        f"Preço inicial inválido para o item {pi.ordem} - "
                        f"{pi.item.nome_item}."
                    )
                    continue

                if preco_inicial <= 0:
                    erros.append(
                        f"O preço inicial do item {pi.ordem} - "
                        f"{pi.item.nome_item} deve ser maior que zero."
                    )
                    continue

                dados_validados.append(
                    {
                        "pregao_item": pi,
                        "participa": True,
                        "marca": marca,
                        "preco_inicial": preco_inicial,
                    }
                )

            if erros:
                for erro in erros:
                    messages.error(request, erro)

                # Mantém na tela exatamente o que o usuário acabou de digitar.
                propostas_digitadas = {}
                for pi in itens_pregao:
                    participa = request.POST.get(
                        f"participa_item_{pi.id}"
                    ) == "on"

                    propostas_digitadas[pi.id] = {
                        "participa": participa,
                        "marca": request.POST.get(
                            f"marca_item_{pi.id}",
                            "",
                        ).strip(),
                        "preco_inicial": request.POST.get(
                            f"preco_item_{pi.id}",
                            "",
                        ).strip(),
                    }

                return render(
                    request,
                    "pregoes/propostas_iniciais.html",
                    {
                        "pregoes": pregoes_lista,
                        "modo": modo,
                        "pregao": pregao,
                        "itens_pregao": itens_pregao,
                        "item_pregao": None,
                        "fornecedores": fornecedores,
                        "fornecedor_selecionado": fornecedor_selecionado,
                        "propostas_salvas": propostas_digitadas,
                        "proposta_digital_import_token": proposta_digital_import_token,
            "pode_controlar_propostas_digitais": pode_controlar_propostas_digitais,
                    },
                )

            salvos = 0
            participantes = 0

            try:
                with transaction.atomic():
                    proposta_digital_bloqueada = None

                    if proposta_digital_para_importar_id:
                        proposta_digital_bloqueada = (
                            PropostaDigital.objects.select_for_update()
                            .filter(
                                id=proposta_digital_para_importar_id,
                                pregao=pregao,
                            )
                            .first()
                        )

                        if (
                            not proposta_digital_bloqueada
                            or proposta_digital_bloqueada.status
                            != PropostaDigital.STATUS_FINALIZADA
                        ):
                            raise ValidationError(
                                "Esta Proposta Digital já foi importada, invalidada "
                                "ou não está mais disponível."
                            )

                        cnpj_fornecedor = _somente_digitos(
                            fornecedor_selecionado.cnpj
                        )
                        cnpj_hash_fornecedor = hashlib.sha256(
                            cnpj_fornecedor.encode("utf-8")
                        ).hexdigest()

                        if (
                            len(cnpj_fornecedor) != 14
                            or not hmac.compare_digest(
                                proposta_digital_bloqueada.cnpj_hash,
                                cnpj_hash_fornecedor,
                            )
                            or not hmac.compare_digest(
                                proposta_digital_bloqueada.estrutura_hash,
                                _hash_estrutura_itens_proposta_digital(pregao),
                            )
                        ):
                            raise ValidationError(
                                "A Proposta Digital não corresponde ao "
                                "pregão/fornecedor selecionado."
                            )

                    for dados in dados_validados:
                        PropostaInicialItem.objects.update_or_create(
                            pregao=pregao,
                            pregao_item=dados["pregao_item"],
                            fornecedor=fornecedor_selecionado,
                            defaults={
                                "participa": dados["participa"],
                                "marca": dados["marca"],
                                "preco_inicial": dados["preco_inicial"],
                            },
                        )
                        salvos += 1
                        if dados["participa"]:
                            participantes += 1

                    if proposta_digital_bloqueada:
                        proposta_digital_bloqueada.status = (
                            PropostaDigital.STATUS_IMPORTADA
                        )
                        proposta_digital_bloqueada.fornecedor_importado = (
                            fornecedor_selecionado
                        )
                        proposta_digital_bloqueada.importado_por = request.user
                        proposta_digital_bloqueada.importado_em = timezone.now()
                        proposta_digital_bloqueada.save(
                            update_fields=[
                                "status",
                                "fornecedor_importado",
                                "importado_por",
                                "importado_em",
                            ]
                        )

            except ValidationError as erro_importacao:
                messages.error(request, str(erro_importacao))
                return redirect(
                    f"{request.path}?modo=fornecedor&pregao={pregao.id}"
                    f"&fornecedor={fornecedor_selecionado.id}"
                )

            mensagem_sucesso = (
                f"Propostas do fornecedor {fornecedor_selecionado.razao_social} "
                f"salvas com sucesso. Itens participantes: {participantes}. "
                f"Registros processados: {salvos}."
            )

            if proposta_digital_para_importar_id:
                mensagem_sucesso += " Proposta Digital marcada como importada."

            messages.success(request, mensagem_sucesso)

            return redirect(
                f"{request.path}?modo=fornecedor&pregao={pregao.id}"
                f"&fornecedor={fornecedor_selecionado.id}"
            )

    return render(
        request,
        "pregoes/propostas_iniciais.html",
        {
            "pregoes": pregoes_lista,
            "modo": modo,
            "pregao": pregao,
            "itens_pregao": itens_pregao,
            "item_pregao": item_pregao,
            "fornecedores": fornecedores,
            "fornecedor_selecionado": fornecedor_selecionado,
            "propostas_salvas": propostas_salvas,
            "proposta_digital_import_token": proposta_digital_import_token,
            "pode_controlar_propostas_digitais": pode_controlar_propostas_digitais,
        },
    )

